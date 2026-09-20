import csv
import hashlib
import io
import json
import re
import tempfile
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO, cast

import jdatetime
from openpyxl import load_workbook

from app.imports.models import IssueSeverity, SourceKind

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
HEADER_CLEANUP = re.compile(r"[^0-9a-zA-Z\u0600-\u06ff]+")
THOUSANDS = re.compile(r"[,٬،\s]")

FIELD_LABELS = {
    "entry_id": "شماره سند",
    "entry_date": "تاریخ سند",
    "account_code": "کد حساب",
    "account_name": "نام حساب",
    "description": "شرح",
    "debit": "بدهکار",
    "credit": "بستانکار",
    "booking_date": "تاریخ تراکنش",
    "amount_signed": "مبلغ خالص",
    "deposit_amount": "واریز",
    "withdrawal_amount": "برداشت",
    "transaction_id": "شناسه تراکنش",
    "transaction_type": "نوع تراکنش",
    "reference": "مرجع",
    "value_date": "تاریخ مؤثر",
    "running_balance": "مانده جاری",
    "counterparty_name": "طرف حساب",
    "iban": "شماره شبا",
    "line_id": "شماره ردیف سند",
    "invoice_ref": "مرجع فاکتور",
    "invoice_no": "شماره فاکتور",
    "issue_date": "تاریخ فاکتور",
    "customer_name": "نام مشتری",
    "gross_amount": "مبلغ ناخالص",
    "due_date": "تاریخ سررسید",
    "tax_amount": "مالیات",
    "paid_amount": "مبلغ وصول‌شده",
    "payment_date": "تاریخ پرداخت",
    "customer_national_id": "شناسه ملی مشتری",
    "status": "وضعیت",
}

ALIASES: dict[str, tuple[str, ...]] = {
    "entry_id": ("شماره سند", "شماره", "سند", "entry id", "document no", "voucher no"),
    "entry_date": ("تاریخ سند", "تاریخ", "entry date", "document date", "date"),
    "account_code": ("کد حساب", "کد", "account code", "ledger code"),
    "account_name": ("نام حساب", "حساب", "account name", "ledger name"),
    "description": (
        "شرح",
        "شرح سند",
        "شرح تراکنش",
        "شرح عملیات",
        "شرح ردیف",
        "توضیحات",
        "description",
        "memo",
        "details",
    ),
    "debit": ("بدهکار", "مبلغ بدهکار", "debit", "debtor"),
    "credit": ("بستانکار", "مبلغ بستانکار", "credit", "creditor"),
    "booking_date": ("تاریخ تراکنش", "تاریخ", "تاریخ عملیات", "booking date", "date"),
    "amount_signed": ("مبلغ", "مبلغ (ریال)", "مبلغ ریال", "مبلغ خالص", "amount", "signed amount"),
    "deposit_amount": (
        "واریز",
        "واریز (ریال)",
        "واریز ریال",
        "واریز(ریال)",
        "مبلغ واریز",
        "بستانکار",
        "بستانکار (ریال)",
        "بستانکار ریال",
        "واریزی",
        "deposit",
        "credit amount",
    ),
    "withdrawal_amount": (
        "برداشت",
        "برداشت (ریال)",
        "برداشت ریال",
        "برداشت(ریال)",
        "مبلغ برداشت",
        "بدهکار",
        "بدهکار (ریال)",
        "بدهکار ریال",
        "برداشتی",
        "withdrawal",
        "debit amount",
    ),
    "transaction_id": (
        "شناسه تراکنش",
        "شماره پیگیری",
        "شماره سند",
        "شماره تراکنش",
        "کد رهگیری",
        "کد پیگیری",
        "شماره ارجاع",
        "شماره فیش",
        "transaction id",
        "tracking id",
    ),
    "transaction_type": ("نوع تراکنش", "نوع عملیات", "نوع", "transaction type", "type"),
    "reference": ("مرجع", "شماره مرجع", "reference", "ref"),
    "value_date": ("تاریخ موثر", "تاریخ ارزش", "value date"),
    "running_balance": (
        "مانده",
        "مانده (ریال)",
        "مانده ریال",
        "مانده(ریال)",
        "مانده جاری",
        "موجودی",
        "موجودی (ریال)",
        "موجودی ریال",
        "تراز",
        "تراز حساب",
        "تراز مانده",
        "running balance",
        "balance",
    ),
    "counterparty_name": ("طرف حساب", "نام طرف حساب", "counterparty", "party name"),
    "iban": ("شبا", "شماره شبا", "iban"),
    "line_id": ("شماره ردیف", "ردیف", "ردیف سند", "ردیف تراکنش", "line id", "row"),
    "invoice_ref": ("مرجع فاکتور", "شماره فاکتور", "invoice ref"),
    "invoice_no": ("شماره فاکتور", "فاکتور", "invoice no", "invoice number"),
    "issue_date": ("تاریخ فاکتور", "تاریخ صدور", "issue date", "invoice date", "date"),
    "customer_name": ("نام مشتری", "مشتری", "طرف حساب", "customer", "customer name"),
    "gross_amount": ("مبلغ ناخالص", "مبلغ کل", "جمع فاکتور", "gross amount", "total"),
    "due_date": ("تاریخ سررسید", "سررسید", "due date"),
    "tax_amount": ("مالیات", "مالیات ارزش افزوده", "tax", "vat"),
    "paid_amount": ("مبلغ وصول شده", "پرداخت شده", "paid amount", "paid"),
    "payment_date": ("تاریخ پرداخت", "payment date"),
    "customer_national_id": ("شناسه ملی مشتری", "کد ملی مشتری", "national id"),
    "status": ("وضعیت", "status"),
}

REQUIRED_FIELDS: dict[SourceKind, tuple[str, ...]] = {
    SourceKind.ACCOUNTING: (
        "entry_date",
        "account_code",
        "account_name",
        "description",
        "debit",
        "credit",
    ),
    SourceKind.BANK: ("booking_date", "description"),
    SourceKind.SALES: ("invoice_no", "issue_date", "customer_name", "gross_amount"),
}

ALTERNATIVE_REQUIRED_FIELDS: dict[SourceKind, tuple[tuple[str, ...], ...]] = {
    SourceKind.ACCOUNTING: (),
    SourceKind.BANK: (("amount_signed",), ("deposit_amount", "withdrawal_amount")),
    SourceKind.SALES: (),
}

TARGET_FIELDS: dict[SourceKind, tuple[str, ...]] = {
    SourceKind.ACCOUNTING: (
        "entry_id",
        "entry_date",
        "account_code",
        "account_name",
        "description",
        "debit",
        "credit",
        "line_id",
        "reference",
        "counterparty_name",
        "invoice_ref",
    ),
    SourceKind.BANK: (
        "booking_date",
        "description",
        "amount_signed",
        "deposit_amount",
        "withdrawal_amount",
        "transaction_id",
        "reference",
        "transaction_type",
        "line_id",
        "value_date",
        "running_balance",
        "counterparty_name",
        "iban",
    ),
    SourceKind.SALES: (
        "invoice_no",
        "issue_date",
        "customer_name",
        "gross_amount",
        "due_date",
        "tax_amount",
        "paid_amount",
        "payment_date",
        "customer_national_id",
        "status",
        "description",
    ),
}

DATE_FIELDS = {
    "entry_date",
    "booking_date",
    "value_date",
    "issue_date",
    "due_date",
    "payment_date",
}
MONEY_FIELDS = {
    "debit",
    "credit",
    "amount_signed",
    "deposit_amount",
    "withdrawal_amount",
    "gross_amount",
    "tax_amount",
    "paid_amount",
    "running_balance",
}


class TableParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedTable:
    sheets: list[str]
    selected_sheet: str
    columns: list[str]
    rows: list[tuple[int, dict[str, object]]]
    header_row: int = 1


@dataclass(frozen=True)
class RowIssue:
    row_number: int | None
    field: str | None
    severity: IssueSeverity
    code: str
    message: str
    raw_value: str | None = None
    remedy: str | None = None


def normalize_digits(value: object) -> str:
    return str(value).translate(PERSIAN_DIGITS)


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", normalize_digits(value)).casefold().strip()
    text = text.replace("ي", "ی").replace("ك", "ک")
    return HEADER_CLEANUP.sub(" ", text).strip()


def column_fingerprint(columns: list[str]) -> str:
    payload = json.dumps([normalize_header(column) for column in columns], ensure_ascii=False)
    return hashlib.sha256(payload.encode()).hexdigest()


def raw_row_hash(raw: dict[str, object]) -> str:
    payload = json.dumps(
        raw, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":")
    )
    return hashlib.sha256(payload.encode()).hexdigest()


def _cell_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _evaluate_header_candidate(
    values: Sequence[object],
) -> tuple[list[str] | None, list[int] | None, int]:
    raw_strings = [str(_cell_value(v)).strip() for v in values]
    all_known_headers = {
        normalize_header(alias)
        for aliases in ALIASES.values()
        for alias in aliases
    } | {normalize_header(label) for label in FIELD_LABELS.values()}

    # Option A: Dense row (all columns from start to end are non-empty)
    dense_cols: list[str] | None = list(raw_strings)
    dense_indices: list[int] | None = None
    dense_score = 0
    while dense_cols and not dense_cols[-1]:
        dense_cols.pop()
    if (
        dense_cols
        and len(dense_cols) >= 2
        and all(dense_cols)
        and len(set(dense_cols)) == len(dense_cols)
    ):
        dense_score = sum(1 for col in dense_cols if normalize_header(col) in all_known_headers)
        dense_indices = list(range(len(dense_cols)))
    else:
        dense_cols, dense_score, dense_indices = None, 0, None

    # Option B: Sparse / Merged cells row (columns at specific non-empty indices)
    non_empty = [(idx, s) for idx, s in enumerate(raw_strings) if s]
    sparse_cols: list[str] | None = None
    sparse_indices: list[int] | None = None
    sparse_score = 0
    if len(non_empty) >= 2:
        sparse_candidates = [s for _, s in non_empty]
        if len(set(sparse_candidates)) == len(sparse_candidates):
            sparse_score = sum(
                1 for col in sparse_candidates if normalize_header(col) in all_known_headers
            )
            sparse_cols = sparse_candidates
            sparse_indices = [idx for idx, _ in non_empty]

    if dense_cols is not None and dense_score >= sparse_score:
        return dense_cols, dense_indices, dense_score
    if sparse_cols is not None and sparse_score >= 2:
        return sparse_cols, sparse_indices, sparse_score
    if dense_cols is not None:
        return dense_cols, dense_indices, dense_score
    if sparse_cols is not None:
        return sparse_cols, sparse_indices, sparse_score
    return None, None, 0


def _columns(values: Sequence[object]) -> tuple[list[str], list[int]]:
    cols, indices, _ = _evaluate_header_candidate(values)
    if cols is None or indices is None:
        raw = [str(_cell_value(v)).strip() for v in values]
        if not any(raw):
            raise TableParseError("ردیف عنوان ستون‌ها خالی است.")
        non_empty = [s for s in raw if s]
        if len(set(non_empty)) != len(non_empty):
            raise TableParseError("عنوان ستون تکراری در فایل وجود دارد.")
        raise TableParseError("ردیف عنوان انتخاب‌شده معتبر نیست یا همه ستون‌ها عنوان ندارند.")
    return cols, indices


def parse_table(
    stream: BinaryIO,
    extension: str,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
    limit: int | None = None,
) -> ParsedTable:
    stream.seek(0)
    if extension == ".csv":
        return _parse_csv(stream, header_row=header_row, limit=limit)
    if extension == ".xlsx":
        return _parse_xlsx(stream, sheet_name=sheet_name, header_row=header_row, limit=limit)
    raise TableParseError("قالب فایل برای پیش‌نمایش پشتیبانی نمی‌شود.")


def _parse_csv(stream: BinaryIO, *, header_row: int, limit: int | None) -> ParsedTable:
    payload = stream.read()
    for encoding in ("utf-8-sig", "cp1256"):
        try:
            text = payload.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TableParseError("رمزگذاری فایل CSV قابل تشخیص نیست.")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    all_lines = list(csv.reader(io.StringIO(text), dialect))
    if not all_lines:
        raise TableParseError("فایل انتخابی خالی است.")

    effective_header_row = header_row
    chosen_cols: list[str] | None = None
    chosen_indices: list[int] | None = None

    if header_row == 1 and len(all_lines) > 1:
        cand_cols, cand_indices, cand_score = _evaluate_header_candidate(all_lines[0])
        if cand_cols is None or cand_score < 2:
            best_row = 1
            best_score = cand_score if cand_cols is not None else 0
            best_cols = cand_cols
            best_indices = cand_indices
            for idx, line in enumerate(all_lines[:25], start=1):
                cols, indices, score = _evaluate_header_candidate(line)
                if cols is not None and score > best_score:
                    best_score = score
                    best_row = idx
                    best_cols = cols
                    best_indices = indices
            if best_score >= 2 and best_cols is not None and best_indices is not None:
                effective_header_row = best_row
                chosen_cols = best_cols
                chosen_indices = best_indices

    if effective_header_row > len(all_lines):
        raise TableParseError("ردیف عنوان انتخاب‌شده در فایل وجود ندارد.")
    header_values = all_lines[effective_header_row - 1]
    if chosen_cols is None or chosen_indices is None:
        chosen_cols, chosen_indices = _columns(header_values)

    columns = chosen_cols
    indices = chosen_indices

    footer_patterns = ("تاریخ صدور", "صفحه ", "از 1", "از ۱", "page ")
    rows: list[tuple[int, dict[str, object]]] = []
    for index, values in enumerate(
        all_lines[effective_header_row:], start=effective_header_row + 1
    ):
        if not any(str(value).strip() for value in values):
            continue
        row_text = " ".join(str(v) for v in values if v)
        if any(p in row_text for p in footer_patterns) and not any(
            str(values[i]).strip()
            for i in indices
            if i < len(values)
            and normalize_header(columns[indices.index(i)]) in ("تاریخ", "date")
        ):
            continue
        row_dict: dict[str, object] = {
            columns[i]: values[col_idx] if col_idx < len(values) else ""
            for i, col_idx in enumerate(indices)
        }
        if not any(str(v).strip() for v in row_dict.values()):
            continue
        rows.append((index, row_dict))
        if limit is not None and len(rows) >= limit:
            break
    return ParsedTable(
        sheets=["CSV"],
        selected_sheet="CSV",
        columns=columns,
        rows=rows,
        header_row=effective_header_row,
    )


def _parse_xlsx(
    stream: BinaryIO,
    *,
    sheet_name: str | None,
    header_row: int,
    limit: int | None,
) -> ParsedTable:
    workbook = load_workbook(stream, read_only=True, data_only=False)
    try:
        sheets = list(workbook.sheetnames)
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise TableParseError("شیت انتخاب‌شده در فایل وجود ندارد.")
        worksheet = workbook[selected]
        if hasattr(worksheet, "reset_dimensions"):
            worksheet.reset_dimensions()
        all_raw_rows: list[tuple[int, list[object]]] = []
        for row_number, cells in enumerate(worksheet.iter_rows(values_only=False), start=1):
            values = [
                str(cell.value)
                if getattr(cell, "data_type", None) == "f"
                else _cell_value(cell.value)
                for cell in cells
            ]
            all_raw_rows.append((row_number, values))

        if not all_raw_rows:
            raise TableParseError("شیت انتخاب‌شده خالی است.")

        effective_header_row = header_row
        chosen_cols: list[str] | None = None
        chosen_indices: list[int] | None = None

        if header_row == 1 and len(all_raw_rows) > 1:
            cand_cols, cand_indices, cand_score = _evaluate_header_candidate(all_raw_rows[0][1])
            if cand_cols is None or cand_score < 2:
                best_row = 1
                best_score = cand_score if cand_cols is not None else 0
                best_cols = cand_cols
                best_indices = cand_indices
                for row_num, line in all_raw_rows[:25]:
                    cols, indices, score = _evaluate_header_candidate(line)
                    if cols is not None and score > best_score:
                        best_score = score
                        best_row = row_num
                        best_cols = cols
                        best_indices = indices
                if best_score >= 2 and best_cols is not None and best_indices is not None:
                    effective_header_row = best_row
                    chosen_cols = best_cols
                    chosen_indices = best_indices

        header_entry = next(
            (item for item in all_raw_rows if item[0] == effective_header_row), None
        )
        if header_entry is None:
            raise TableParseError("ردیف عنوان انتخاب‌شده در فایل وجود ندارد.")

        if chosen_cols is None or chosen_indices is None:
            chosen_cols, chosen_indices = _columns(header_entry[1])

        columns = chosen_cols
        indices = chosen_indices

        footer_patterns = ("تاریخ صدور", "صفحه ", "از 1", "از ۱", "page ")
        rows: list[tuple[int, dict[str, object]]] = []
        for row_number, values in all_raw_rows:
            if row_number <= effective_header_row:
                continue
            if not any(str(value).strip() for value in values):
                continue
            row_text = " ".join(str(v) for v in values if v)
            if any(p in row_text for p in footer_patterns) and not any(
                str(values[i]).strip()
                for i in indices
                if i < len(values)
                and normalize_header(columns[indices.index(i)]) in ("تاریخ", "date")
            ):
                continue
            row_dict = {
                columns[i]: values[col_idx] if col_idx < len(values) else ""
                for i, col_idx in enumerate(indices)
            }
            if not any(str(v).strip() for v in row_dict.values()):
                continue
            rows.append((row_number, row_dict))
            if limit is not None and len(rows) >= limit:
                break
        return ParsedTable(
            sheets=sheets,
            selected_sheet=selected,
            columns=columns,
            rows=rows,
            header_row=effective_header_row,
        )
    finally:
        workbook.close()


def download_to_seekable(body: Any) -> BinaryIO:
    stream = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        while chunk := body.read(1024 * 1024):
            stream.write(chunk)
        stream.seek(0)
        return cast(BinaryIO, stream)
    except Exception:
        stream.close()
        raise
    finally:
        body.close()


def suggest_mapping(columns: list[str], source_kind: SourceKind) -> list[dict[str, object]]:
    normalized_columns = {normalize_header(column): column for column in columns}
    targets = list(TARGET_FIELDS[source_kind])
    suggestions: list[dict[str, object]] = []
    used: set[str] = set()
    for target in targets:
        aliases = (FIELD_LABELS.get(target, target), *ALIASES.get(target, ()))
        for alias in aliases:
            source = normalized_columns.get(normalize_header(alias))
            if source is not None and source not in used:
                suggestions.append(
                    {
                        "target_field": target,
                        "source_column": source,
                        "confidence": 100
                        if normalize_header(alias) == normalize_header(target)
                        else 92,
                        "reason": "تطابق عنوان ستون با فرهنگ نام‌های فارسی و انگلیسی",
                    }
                )
                used.add(source)
                break
    return suggestions


def validate_mapping_fields(
    source_kind: SourceKind, mapping: dict[str, str], columns: list[str]
) -> list[str]:
    errors: list[str] = []
    unknown_columns = sorted(set(mapping.values()) - set(columns))
    if unknown_columns:
        errors.append(f"ستون‌های منبع نامعتبر: {', '.join(unknown_columns)}")
    missing = [field for field in REQUIRED_FIELDS[source_kind] if field not in mapping]
    if missing:
        errors.append(f"فیلدهای الزامی نگاشت نشده‌اند: {', '.join(missing)}")
    alternatives = ALTERNATIVE_REQUIRED_FIELDS[source_kind]
    if alternatives and not any(all(field in mapping for field in group) for group in alternatives):
        readable = " یا ".join(" + ".join(group) for group in alternatives)
        errors.append(f"یکی از ترکیب‌های مبلغ الزامی است: {readable}")
    allowed = set(TARGET_FIELDS[source_kind])
    unknown_targets = sorted(set(mapping) - allowed)
    if unknown_targets:
        errors.append(f"فیلدهای مقصد ناشناخته‌اند: {', '.join(unknown_targets)}")
    return errors


def validate_transforms(
    mapping: Mapping[str, str], transforms: Mapping[str, Sequence[str]], currency_unit: str
) -> list[str]:
    errors: list[str] = []
    unknown_fields = sorted(set(transforms) - set(mapping))
    if unknown_fields:
        errors.append(f"تبدیل برای فیلدهای نگاشت‌نشده تعریف شده است: {', '.join(unknown_fields)}")
    for field in mapping:
        operations = set(transforms.get(field, []))
        if field in DATE_FIELDS:
            missing = {"normalize_digits", "parse_date"} - operations
            if missing:
                errors.append(
                    f"تبدیل تاریخ برای {field} باید صریحاً شامل {', '.join(sorted(missing))} باشد."
                )
        if field in MONEY_FIELDS:
            missing = {"normalize_digits", "strip_thousands"} - operations
            if missing:
                errors.append(
                    f"تبدیل مبلغ برای {field} باید صریحاً شامل {', '.join(sorted(missing))} باشد."
                )
            if currency_unit == "toman" and "toman_to_rial" not in operations:
                errors.append(f"تبدیل تومان به ریال برای {field} باید صریحاً تأیید شود.")
            if currency_unit == "rial" and "toman_to_rial" in operations:
                errors.append(f"تبدیل تومان به ریال با واحد ریال برای {field} ناسازگار است.")
    return errors


def _parse_date(value: object, calendar: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    normalized = normalize_digits(value).strip().replace(".", "/").replace("-", "/")
    date_part = normalized.split()[0].split("t")[0].split("T")[0]
    parts = date_part.split("/")
    if len(parts) != 3:
        raise ValueError("invalid date")
    year, month, day = (int(part) for part in parts)
    if calendar == "jalali":
        return str(jdatetime.date(year, month, day).togregorian().isoformat())
    return date(year, month, day).isoformat()


def _parse_money(value: object, *, toman: bool) -> str:
    normalized = THOUSANDS.sub("", normalize_digits(value)).strip()
    if normalized in ("", "-", "–", "—", "None", "null", "."):
        return "0"
    try:
        amount = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError("invalid money") from exc
    if not amount.is_finite():
        raise ValueError("invalid money")
    if toman:
        amount *= 10
    if amount != amount.to_integral_value():
        raise ValueError("fractional rial")
    return str(int(amount))


def transform_and_validate_row(
    *,
    source_kind: SourceKind,
    row_number: int,
    raw: dict[str, object],
    mapping: dict[str, str],
    transforms: dict[str, list[str]],
    currency_unit: str,
    calendar: str,
) -> tuple[dict[str, object], list[RowIssue]]:
    transformed: dict[str, object] = {}
    issues: list[RowIssue] = []
    for target, source in mapping.items():
        raw_value = raw.get(source, "")
        if isinstance(raw_value, str) and raw_value.startswith("="):
            issues.append(
                RowIssue(
                    row_number,
                    target,
                    IssueSeverity.ERROR,
                    "FORMULA_NOT_ALLOWED",
                    "فرمول اکسل به‌عنوان داده پذیرفته نمی‌شود.",
                    raw_value,
                    "فرمول را در فایل منبع به مقدار ثابت تبدیل کنید.",
                )
            )
            continue
        value: object = raw_value
        operations = transforms.get(target, [])
        try:
            if "trim" in operations and isinstance(value, str):
                value = value.strip()
            if "normalize_digits" in operations:
                value = normalize_digits(value)
            if target in DATE_FIELDS or "parse_date" in operations:
                value = _parse_date(value, calendar)
            if target in MONEY_FIELDS or "strip_thousands" in operations:
                value = _parse_money(value, toman="toman_to_rial" in operations)
            if "withdrawal_negative" in operations and Decimal(str(value)) > 0:
                value = str(-Decimal(str(value)))
            transformed[target] = value
        except (ValueError, InvalidOperation):
            code = "INVALID_DATE" if target in DATE_FIELDS else "INVALID_MONEY"
            label = "تاریخ" if target in DATE_FIELDS else "مبلغ"
            issues.append(
                RowIssue(
                    row_number,
                    target,
                    IssueSeverity.ERROR,
                    code,
                    f"{label} قابل خواندن نیست.",
                    str(raw_value)[:500],
                    f"قالب {label} و تقویم/واحد انتخاب‌شده را بررسی کنید.",
                )
            )

    required = REQUIRED_FIELDS[source_kind]
    for field in required:
        if str(transformed.get(field, "")).strip() == "":
            issues.append(
                RowIssue(
                    row_number,
                    field,
                    IssueSeverity.ERROR,
                    "REQUIRED_VALUE_MISSING",
                    f"مقدار الزامی «{FIELD_LABELS.get(field, field)}» خالی است.",
                    None,
                    "مقدار را در فایل منبع تکمیل کنید.",
                )
            )

    if source_kind == SourceKind.ACCOUNTING:
        debit = Decimal(str(transformed.get("debit", "0")))
        credit = Decimal(str(transformed.get("credit", "0")))
        if debit < 0 or credit < 0:
            issues.append(
                RowIssue(
                    row_number,
                    None,
                    IssueSeverity.ERROR,
                    "NEGATIVE_LEDGER_AMOUNT",
                    "بدهکار و بستانکار نمی‌توانند منفی باشند.",
                )
            )
        if debit > 0 and credit > 0:
            issues.append(
                RowIssue(
                    row_number,
                    None,
                    IssueSeverity.ERROR,
                    "BOTH_DEBIT_AND_CREDIT",
                    "یک ردیف نمی‌تواند هم‌زمان بدهکار و بستانکار باشد.",
                )
            )
        if "entry_id" not in mapping:
            issues.append(
                RowIssue(
                    row_number,
                    "entry_id",
                    IssueSeverity.WARNING,
                    "ENTRY_ID_MISSING",
                    "شماره سند نگاشت نشده و کنترل تراز سند محدود است.",
                    remedy="برای کنترل تراز، شماره سند را نگاشت کنید.",
                )
            )
    elif source_kind == SourceKind.BANK and "amount_signed" not in mapping:
        deposit = Decimal(str(transformed.get("deposit_amount", "0")))
        withdrawal = Decimal(str(transformed.get("withdrawal_amount", "0")))
        if deposit < 0 or withdrawal < 0:
            issues.append(
                RowIssue(
                    row_number,
                    None,
                    IssueSeverity.ERROR,
                    "NEGATIVE_BANK_COMPONENT",
                    "مبلغ واریز و برداشت در ستون‌های جداگانه نمی‌تواند منفی باشد.",
                )
            )
        if deposit > 0 and withdrawal > 0:
            issues.append(
                RowIssue(
                    row_number,
                    None,
                    IssueSeverity.ERROR,
                    "BOTH_DEPOSIT_AND_WITHDRAWAL",
                    "یک تراکنش نمی‌تواند هم‌زمان واریز و برداشت باشد.",
                )
            )
        transformed["amount_signed"] = str(deposit - withdrawal)
    elif source_kind == SourceKind.SALES and "due_date" not in mapping:
        issues.append(
            RowIssue(
                row_number,
                "due_date",
                IssueSeverity.WARNING,
                "DUE_DATE_MISSING",
                "تاریخ سررسید موجود نیست و تحلیل وصول محدود می‌شود.",
                remedy="در صورت وجود، ستون سررسید را نگاشت کنید.",
            )
        )
    return transformed, issues


def coverage_for(
    source_kind: SourceKind,
    *,
    row_count: int,
    accepted_count: int,
    mapping: dict[str, str],
) -> dict[str, object]:
    overall = round(100 * accepted_count / row_count) if row_count else 0
    coverage: dict[str, object] = {
        "overall": overall,
        "rows": {
            "total": row_count,
            "accepted": accepted_count,
            "rejected": row_count - accepted_count,
        },
    }
    if source_kind == SourceKind.ACCOUNTING:
        coverage["journal_balance"] = {
            "available": "entry_id" in mapping,
            "score": overall if "entry_id" in mapping else 0,
            "reasons": [] if "entry_id" in mapping else ["شماره سند نگاشت نشده است"],
        }
    elif source_kind == SourceKind.BANK:
        coverage["reconciliation"] = {
            "available": accepted_count > 0,
            "score": overall,
            "reasons": [] if accepted_count else ["تراکنش معتبر وجود ندارد"],
        }
    else:
        coverage["receivables"] = {
            "available": "due_date" in mapping and accepted_count > 0,
            "score": overall if "due_date" in mapping else 0,
            "reasons": [] if "due_date" in mapping else ["تاریخ سررسید نگاشت نشده است"],
        }
    return coverage
