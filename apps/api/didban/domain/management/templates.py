"""A bounded, explicit Excel mapping format; workbook formulas never execute."""

import base64
from io import BytesIO
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from .statements import SECTIONS


class TemplateError(ValueError):
    pass


HEADERS = ("کلید", "عنوان", "بخش", "حساب‌ها")


def validate_template(template: dict) -> dict:
    if (
        not template
        or not template.get("name", "").strip()
        or len(template["name"]) > 120
    ):
        raise TemplateError("نام قالب را وارد کنید.")
    lines = template.get("lines", [])
    if not 1 <= len(lines) <= 100:
        raise TemplateError("قالب باید بین ۱ تا ۱۰۰ سرفصل داشته باشد.")
    keys, accounts = set(), set()
    for line in lines:
        if (
            not line["key"].strip()
            or not line["label"].strip()
            or max(len(line["key"]), len(line["label"])) > 120
            or line["key"] in keys
        ):
            raise TemplateError("کلید و عنوان سرفصل‌ها باید مشخص و کلیدها یکتا باشند.")
        keys.add(line["key"])
        if line["section"] not in SECTIONS:
            raise TemplateError("بخش سرفصل معتبر نیست.")
        if len(line["accounts"]) != len(set(line["accounts"])):
            raise TemplateError("یک حساب در یک سرفصل تکرار شده است.")
        for account in line["accounts"]:
            if (
                not account.strip()
                or len(account) > 120
                or account != account.strip()
                or account in accounts
            ):
                raise TemplateError("هر حساب فقط باید در یک سرفصل قرار بگیرد.")
            accounts.add(account)
    if len(accounts) > 2000:
        raise TemplateError("قالب حداکثر ۲۰۰۰ حساب می‌پذیرد.")
    return template


def parse_template(name: str, encoded: str) -> dict:
    try:
        data = base64.b64decode(encoded, validate=True)
        if len(data) > 30000:
            raise TemplateError("حجم قالب باید کمتر از ۳۰ کیلوبایت باشد.")
        with ZipFile(BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 100
                or sum(i.file_size for i in archive.infolist()) > 1_000_000
            ):
                raise TemplateError("ساختار فایل قالب بیش از حد بزرگ است.")
            if any(
                "vbaProject" in i.filename or "externalLinks/" in i.filename
                for i in archive.infolist()
            ):
                raise TemplateError("قالب نباید ماکرو یا پیوند خارجی داشته باشد.")
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if workbook.sheetnames != ["نگاشت"]:
                raise TemplateError("فایل باید فقط برگه «نگاشت» داشته باشد.")
            sheet = workbook["نگاشت"]
            if sheet.max_row > 101 or sheet.max_column != 4:
                raise TemplateError("ساختار سطرها و ستون‌های قالب معتبر نیست.")
            if (
                tuple(cell.value for cell in next(sheet.iter_rows(max_row=1)))
                != HEADERS
            ):
                raise TemplateError(
                    "ستون‌های قالب باید «کلید، عنوان، بخش، حساب‌ها» باشند."
                )
            lines = []
            for row in sheet.iter_rows(min_row=2, max_col=4):
                if all(c.value is None for c in row):
                    continue
                if any(
                    c.data_type == "f"
                    or (c.value is not None and not isinstance(c.value, str))
                    for c in row
                ):
                    raise TemplateError(
                        "مقادیر قالب باید متن باشند؛ کد حساب را متنی ذخیره کنید و فرمول ننویسید."
                    )
                key, label, section, raw = [str(c.value or "").strip() for c in row]
                if max(len(key), len(label)) > 120 or len(raw) > 10000:
                    raise TemplateError("متن یکی از سرفصل‌ها بیش از حد طولانی است.")
                lines.append(
                    {
                        "key": key,
                        "label": label,
                        "section": section,
                        "accounts": [a.strip() for a in raw.split(",") if a.strip()],
                    }
                )
            return validate_template({"name": name.strip(), "lines": lines})
        finally:
            workbook.close()
    except TemplateError:
        raise
    except Exception as exc:
        raise TemplateError("فایل قالب Excel معتبر نیست.") from exc


def template_bytes(template: dict | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "نگاشت"
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS)
    lines = (
        template["lines"]
        if template
        else [
            {"key": key, "label": title, "section": key, "accounts": []}
            for key, (title, _, _) in SECTIONS.items()
        ]
    )
    for line in lines:
        sheet.append(
            [line["key"], line["label"], line["section"], ",".join(line["accounts"])]
        )
    for column in ("A", "B", "C", "D"):
        sheet.column_dimensions[column].width = 30
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def template_from_report(management: dict, rows: list[dict]) -> dict:
    """Include prior-only accounts so the exported mapping can compare both periods."""
    from .statements import ROLE_SECTION

    if management.get("template"):
        return management["template"]
    accounts = set(management["statements"]["available_accounts"])
    groups = {}
    for row in rows:
        values = row["canonical_values"]
        if (
            row["source_kind"] == management.get("statement_basis", "ledger")
            and values.get("account") in accounts
        ):
            groups.setdefault(ROLE_SECTION.get(values.get("role")), set()).add(
                values["account"]
            )
    return {
        "name": "قالب پایه",
        "lines": [
            {key: line[key] for key in ("key", "label", "section")}
            | {"accounts": sorted(groups.get(line["section"], set()))}
            for line in management["statements"]["lines"]
        ],
    }
