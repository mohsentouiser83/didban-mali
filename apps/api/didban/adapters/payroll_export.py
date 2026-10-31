"""Private payroll previews and approved downloads: shaped embedded-font PDF and text-exact Excel."""

from __future__ import annotations
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from datetime import date
from .money_presentation import WorkbookMoney, workbook_money, grouped_amount

LABELS = {
    "gross": "ناخالص",
    "net": "خالص قابل پرداخت",
    "employee_insurance": "بیمه سهم کارمند",
    "employer_insurance": "بیمه سهم کارفرما",
    "tax": "مالیات",
    "deductions": "سایر کسورات",
    "employer_cost": "هزینه کارفرما",
    "leave_balance_days": "مانده مرخصی",
}


def export_run(run, format, settings):
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        book = Workbook()
        sheet = book.active
        sheet.title = "لیست حقوق"
        sheet.sheet_view.rightToLeft = True
        sheet.append(["کد پرسنلی", "شرکت", "نام", *LABELS.values()])
        for item in run["content"]["items"]:
            sheet.append(
                [
                    item["personnel_code"],
                    run.get("company_name", "—"),
                    item["profile"]["name"],
                    *[
                        item["result"][k]
                        if k == "leave_balance_days"
                        else workbook_money(item["result"][k]).cell_value()
                        for k in LABELS
                    ],
                ]
            )
        for row in sheet:
            for cell in row:
                cell.font = Font(name="IRANYekanX", size=11)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for row, item in zip(sheet.iter_rows(min_row=2), run["content"]["items"]):
            for cell, key in zip(row[3:], LABELS):
                if key != "leave_balance_days":
                    cell.number_format = WorkbookMoney(
                        item["result"][key]
                    ).number_format
        metadata = book.create_sheet("مشخصات")
        metadata.sheet_view.rightToLeft = True
        for label, value in [
            ("شرکت", run.get("company_name", "—")),
            ("وضعیت", "تأییدشده" if run.get("approved") else "پیش‌نمایش؛ تأیید نشده"),
            ("شروع بازه (میلادی)", str(run["period_start"])),
            ("پایان بازه (میلادی)", str(run["period_end"])),
            ("شناسه نسخه", str(run["id"])),
            ("واحد مبلغ", "ریال"),
        ]:
            metadata.append([label, value])
        for row in metadata:
            for cell in row:
                cell.font = Font(name="IRANYekanX", size=11)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        out = BytesIO()
        book.save(out)
        return (
            out.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    import arabic_reshaper
    from bidi.algorithm import get_display

    root = Path(__file__).resolve().parents[4]
    fonts = root / "assets/fonts"
    if not fonts.is_dir():
        fonts = root / "apps/web/public/fonts"
    for name, filename in [
        ("PayrollRegular", "IRANYekanX-Regular.ttf"),
        ("PayrollMedium", "IRANYekanX-Medium.ttf"),
    ]:
        path = fonts / filename
        if not path.is_file():
            matches = list(root.glob("**/" + filename))
            path = matches[0] if matches else path
        if not path.is_file():
            raise RuntimeError("payroll_font_missing")
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))

    def text(value):
        return escape(
            get_display(
                arabic_reshaper.reshape(
                    str(value).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))
                )
            )
        )

    style = ParagraphStyle(
        "Payroll",
        fontName="PayrollRegular",
        fontSize=11,
        leading=21,
        alignment=TA_RIGHT,
    )
    title = ParagraphStyle(
        "PayrollTitle", parent=style, fontName="PayrollMedium", fontSize=18, leading=28
    )
    import jdatetime

    def report_date(value):
        return jdatetime.date.fromgregorian(
            date=date.fromisoformat(str(value))
        ).strftime("%Y/%m/%d")

    def money(value):
        return grouped_amount(str(value)).replace(".", "٫")

    story = []
    for i, item in enumerate(run["content"]["items"]):
        if i:
            story.append(PageBreak())
        for value in [
            "فیش حقوقی تأییدشده" if run.get("approved") else "پیش‌نمایش محاسبه حقوق",
            "شرکت: " + run.get("company_name", "—"),
            item["profile"]["name"],
            "کد پرسنلی: " + item["personnel_code"],
            report_date(run["period_start"]) + " تا " + report_date(run["period_end"]),
            "واحد: ریال",
        ]:
            story.append(
                Paragraph(
                    text(value),
                    title if value.startswith(("فیش حقوقی", "پیش‌نمایش")) else style,
                )
            )
        story.append(Spacer(1, 8 * mm))
        line_labels = {
            "base": "مزد مبنا",
            "overtime": "اضافه‌کاری",
            "shift": "نوبت‌کاری",
            "night": "شب‌کاری",
            "holiday": "تعطیل‌کاری",
            "eidi": "عیدی",
            "sanavat": "سنوات",
            "termination": "تسویه",
            "arrears": "معوقات",
            "unpaid_wages_carry": "حقوق قبلی پرداخت‌نشده",
        }
        for line in item["result"]["lines"]:
            story.append(
                Paragraph(
                    text(
                        line_labels.get(line["key"], line["key"])
                        + ": "
                        + money(line["amount"])
                    ),
                    style,
                )
            )
        story.append(Spacer(1, 4 * mm))
        for key, label in LABELS.items():
            story.append(
                Paragraph(text(label + ": " + money(item["result"][key])), style)
            )
        story.append(Paragraph(text("شناسه نسخه محاسبه"), style))
        story.append(
            Paragraph(
                escape(str(run["id"])),
                ParagraphStyle(
                    "PayrollIdentifier", parent=style, alignment=0, fontSize=8
                ),
            )
        )
    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )
    doc.build(story)
    return out.getvalue(), "application/pdf", "pdf"
