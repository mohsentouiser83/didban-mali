"""Bounded employee/work imports; formulas and executable workbook features rejected."""

from __future__ import annotations
import base64
import csv
from io import BytesIO, StringIO
from pathlib import Path
import zipfile
from datetime import date, datetime
from itertools import islice
from uuid import UUID
from ..domain.payroll.validation import employee_profile, work_content, valid_date
from ..domain.payroll.calculation import PayrollValidation, number
from ..domain.financial.common import decimal_text
from ..domain.settlements import exact_context


@exact_context
def parse_import(body, resolve_employee=None):
    try:
        data = base64.b64decode(body["data_base64"], validate=True)
    except (ValueError, TypeError) as exc:
        raise PayrollValidation("invalid_import_encoding") from exc
    if len(data) > 2_000_000:
        raise PayrollValidation("import_too_large")
    suffix = Path(body["filename"]).suffix.lower()
    if suffix == ".csv":
        try:
            reader = csv.DictReader(StringIO(data.decode("utf-8-sig")))
            if (
                not reader.fieldnames
                or len(reader.fieldnames) != len(set(reader.fieldnames))
                or len(reader.fieldnames) > 128
                or any(h.startswith("__") for h in reader.fieldnames)
            ):
                raise PayrollValidation("invalid_import_headers")
            rows = list(islice(reader, 1001))
        except UnicodeError as exc:
            raise PayrollValidation("import_utf8_required") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20_000_000 or any(
                    "vbaProject" in i.filename or "externalLinks/" in i.filename
                    for i in archive.infolist()
                ):
                    raise PayrollValidation("unsafe_workbook")
            book = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            if len(book.worksheets) != 1:
                raise PayrollValidation("one_sheet_required")
            sheet = book.worksheets[0]
            if (sheet.max_row or 0) > 1001 or (sheet.max_column or 0) > 128:
                raise PayrollValidation("import_row_or_column_limit")
            cells = list(
                islice(sheet.iter_rows(max_col=min(sheet.max_column or 128, 128)), 1002)
            )
            if any(c.data_type == "f" for row in cells for c in row):
                raise PayrollValidation("formulas_not_supported")
            if not cells:
                raise PayrollValidation("empty_import")
            headers = [str(c.value or "").strip() for c in cells[0]]
            if (
                not all(headers)
                or len(headers) != len(set(headers))
                or any(h.startswith("__") for h in headers)
            ):
                raise PayrollValidation("invalid_import_headers")
            rows = []
            for source_row, row in enumerate(cells[1:], 2):
                if not any(c.value is not None for c in row):
                    continue
                record = {
                    k: c.value.date().isoformat()
                    if isinstance(c.value, datetime)
                    else c.value.isoformat()
                    if isinstance(c.value, date)
                    else str(c.value)
                    if c.value is not None
                    else ""
                    for k, c in zip(headers, row)
                }
                invalid_identifiers = [
                    k
                    for k, c in zip(headers, row)
                    if k
                    in (
                        "personnel_code",
                        "national_id",
                        "insurance_number",
                        "iban",
                        "employee_id",
                    )
                    and c.value is not None
                    and not isinstance(c.value, str)
                ]
                if invalid_identifiers:
                    record["__invalid_identifier_cell"] = " / ".join(
                        invalid_identifiers
                    )
                record["__source_row"] = str(source_row)
                rows.append(record)
            book.close()
        except PayrollValidation:
            raise
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise PayrollValidation("invalid_workbook") from exc
    else:
        raise PayrollValidation("unsupported_import_format")
    if not rows or len(rows) > 1000:
        raise PayrollValidation("import_row_limit")
    normalized = []
    preview_rows = []
    errors = []
    seen = set()
    resolved_employees = {}
    for i, row in enumerate(rows, 2):
        i = int(row.pop("__source_row", i))
        original = {
            str(k): v
            for k, v in row.items()
            if k is not None and not k.startswith("__")
        }
        column = "*"
        converted = None
        try:
            invalid_identifier = row.pop("__invalid_identifier_cell", None)
            if invalid_identifier:
                column = invalid_identifier
                raise PayrollValidation("identifier_must_be_text")
            if body["kind"] == "employees":
                column = "personnel_code"
                code = row.get("personnel_code", "").strip()
                if not code or code in seen:
                    raise PayrollValidation("duplicate_or_missing_personnel_code")
                seen.add(code)
                profile = {
                    k: v for k, v in row.items() if k != "personnel_code" and v != ""
                }
                monetary = ("opening_loan", "opening_advance", "opening_unpaid_wages")
                column = "unit"
                if profile.get("unit") and profile["unit"] not in ("IRR", "IRT"):
                    raise PayrollValidation("unsupported_payroll_unit")
                if any(profile.get(k) not in (None, "") for k in monetary):
                    unit = profile.pop("unit", None)
                    if unit not in ("IRR", "IRT"):
                        raise PayrollValidation("opening_money_unit_required")
                    for key in monetary:
                        if key in profile:
                            column = key
                            profile[key] = decimal_text(
                                number(profile[key]) * (10 if unit == "IRT" else 1)
                            )
                    profile["opening_unit_original"] = unit
                else:
                    profile.pop("unit", None)
                for key in ("hired_on", "ended_on", "opening_as_of"):
                    if key in profile:
                        column = key
                        valid_date(profile[key])
                for key in ("opening_leave_days",):
                    if key in profile:
                        column = key
                        number(profile[key])
                column = "name" if not profile.get("name", "").strip() else "*"
                if not profile.get("hired_on"):
                    column = "hired_on"
                if profile.get("national_id") and (
                    len(profile["national_id"]) != 10
                    or not profile["national_id"].isascii()
                    or not profile["national_id"].isdigit()
                ):
                    column = "national_id"
                converted = {
                    "personnel_code": code,
                    "profile": employee_profile(profile),
                }
                normalized.append(converted)
            else:
                column = "employee_id"
                employee = row.get("employee_id", "").strip()
                personnel_code = row.get("personnel_code", "").strip()
                start = row.get("period_start", "")
                end = row.get("period_end", "")
                column = "period_start"
                valid_date(start)
                column = "period_end"
                if valid_date(end) < valid_date(start):
                    raise PayrollValidation("invalid_work_period")
                column = "employee_id"
                if employee:
                    try:
                        employee = str(UUID(employee))
                    except (ValueError, TypeError) as exc:
                        raise PayrollValidation("invalid_employee_id") from exc
                if not employee and not personnel_code:
                    raise PayrollValidation("employee_identity_required")
                if resolve_employee:
                    identity = (employee, personnel_code)
                    if identity not in resolved_employees:
                        resolved_employees[identity] = resolve_employee(
                            employee, personnel_code
                        )
                    resolved = resolved_employees[identity]
                    employee = resolved["employee_id"]
                elif personnel_code:
                    raise PayrollValidation("employee_resolution_required")
                else:
                    resolved = {"employee_id": employee}
                if (employee, start, end) in seen:
                    raise PayrollValidation("duplicate_or_missing_work_identity")
                seen.add((employee, start, end))
                content = {
                    k: v
                    for k, v in row.items()
                    if k
                    not in (
                        "employee_id",
                        "personnel_code",
                        "employee_name",
                        "name",
                        "period_start",
                        "period_end",
                    )
                    and v != ""
                }
                monetary = (
                    "deductions",
                    "loan_installment",
                    "advance_offset",
                    "prior_paid_special",
                    "unpaid_wages",
                    "arrears_amount",
                    "prior_taxable",
                    "prior_tax_paid",
                    *(
                        kind + "_" + suffix
                        for kind in ("eidi", "sanavat", "termination", "arrears")
                        for suffix in ("prior_taxable", "prior_tax_paid")
                    ),
                )
                column = "unit"
                unit = content.pop("unit", None)
                if unit is not None and unit not in ("IRR", "IRT"):
                    raise PayrollValidation("unsupported_payroll_unit")
                if any(k in content for k in monetary):
                    if unit not in ("IRR", "IRT"):
                        raise PayrollValidation("work_money_unit_required")
                    for key in monetary:
                        if key in content:
                            column = key
                            content[key] = decimal_text(
                                number(content[key]) * (10 if unit == "IRT" else 1)
                            )
                    content["source_unit"] = unit
                for key in (
                    "tax_history_complete",
                    "unpaid_wages_already_processed",
                    *(
                        kind + "_tax_history_complete"
                        for kind in ("eidi", "sanavat", "termination", "arrears")
                    ),
                ):
                    if key in content:
                        column = key
                        if content[key].lower() not in ("true", "false"):
                            raise PayrollValidation("boolean_true_false_required")
                        content[key] = content[key].lower() == "true"
                for key in ("period_days", "worked_days", "worked_hours"):
                    column = key
                    number(content.get(key))
                for key in (
                    "overtime_hours",
                    "night_hours",
                    "holiday_hours",
                    "leave_opening_days",
                    "leave_accrued_days",
                    "leave_used_days",
                    "leave_cashout_days",
                    "service_days",
                ):
                    if key in content:
                        column = key
                        number(content[key])
                column = (
                    "tax_elapsed_months" if "tax_elapsed_months" in content else "*"
                )
                converted = {
                    **resolved,
                    "employee_id": employee,
                    "period_start": start,
                    "period_end": end,
                    "content": work_content(content),
                }
                normalized.append(converted)
        except PayrollValidation as exc:
            code = str(exc)
            column = {
                "invalid_work_days": "worked_days / period_days",
                "opening_balance_source_required": "opening_as_of / opening_source",
                "invalid_employment_dates": "ended_on",
                "invalid_tax_months": "tax_elapsed_months",
                "employee_id_not_found": "employee_id",
                "personnel_code_not_found": "personnel_code",
                "employee_identity_mismatch": "employee_id / personnel_code",
                "employee_identity_required": "personnel_code / employee_id",
                "employee_resolution_required": "personnel_code",
                "duplicate_or_missing_work_identity": "personnel_code / employee_id / period_start / period_end",
            }.get(code, column)
            errors.append({"row": i, "column": column, "code": code})
        preview_rows.append({"row": i, "original": original, "normalized": converted})
    return data, {
        "kind": body["kind"],
        "rows": normalized,
        "errors": errors,
        "row_count": len(rows),
        "preview_rows": preview_rows,
    }
