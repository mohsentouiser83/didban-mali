"""Safe exact tabular exports for the installed offline application."""

from __future__ import annotations
import csv
from io import BytesIO, StringIO
import json
from ..domain.reports import METRIC_LABELS, spreadsheet_text, display_amount


def csv_bytes(content: dict) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(
        [
            "record_type",
            "report_id",
            "report_hash",
            "source_file_id",
            "sheet",
            "row_number",
            "state",
            "errors",
            "original_values",
        ]
    )
    identity = content.get("_report", {})
    writer.writerow(
        [
            "report",
            identity.get("id", ""),
            identity.get("content_hash", ""),
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    for row in content["rows"]:
        if row["state"] != "valid" or row["error_codes"]:
            writer.writerow(
                [
                    spreadsheet_text(v)
                    for v in [
                        "error",
                        identity.get("id", ""),
                        identity.get("content_hash", ""),
                        row["source_file_id"],
                        row["sheet_name"],
                        row["row_number"],
                        row["state"],
                        ";".join(row["error_codes"]),
                        json.dumps(row["raw_values"], ensure_ascii=False),
                    ]
                ]
            )
    return b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")


def xlsx_bytes(content: dict) -> bytes:
    # This is the installed application's backend exporter, not an agent-created
    # one-off workbook. It runs offline with the existing locked openpyxl runtime.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        ws = book.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A2"
        ws.append(["قطعه سلول", *headers])
        for values in rows:
            texts = ["" if v is None else str(v) for v in values]
            pieces = max([1, *[(len(v) + 29999) // 30000 for v in texts]])
            for part in range(pieces):
                ws.append(
                    [
                        str(part + 1),
                        *[
                            v[part * 30000 : (part + 1) * 30000]
                            if len(v) > 30000
                            else v
                            for v in texts
                        ],
                    ]
                )
        for row in ws:
            for cell in row:
                cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(name="IRANYekanX", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5556DE")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="IRANYekanX", size=11)
                cell.alignment = Alignment(
                    horizontal="right", vertical="top", wrap_text=True
                )
                # Exact decimal strings are text, never Excel's 15-digit floats.
                cell.number_format = "@"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, len(headers) + 2):
            ws.column_dimensions[get_column_letter(i)].width = (
                28 if i < len(headers) else 60
            )

    manifest = content["manifest"]
    unit = manifest.get("display_unit", "IRR")
    unit_fa = "تومان" if unit == "IRT" else "ریال"
    report = content.get("_report", {})
    sheet(
        "خلاصه",
        ["عنوان", "مقدار"],
        [
            ["برچسب", content["label"]],
            ["شرکت", manifest["company_name"]],
            ["دوره", manifest["period_name"]],
            ["تاریخ مبنا", manifest["as_of_date"]],
            ["شناسه نسخه", manifest["id"]],
            ["هش داده", manifest["manifest_hash"]],
            ["شناسه گزارش", report.get("id", "")],
            ["هش گزارش", report.get("content_hash", "")],
            ["نسخه گزارش", report.get("revision_number", "")],
            ["واحد شاخص‌ها و پیش‌بینی", unit],
            ["واحد پایه داده‌های استاندارد", "IRR"],
            ["مبالغ", "رشته دقیق؛ ناموجود با وضعیت و علت مشخص است"],
        ],
    )
    sheet(
        "منابع",
        ["شناسه فایل", "نام اصلی", "هش اصل فایل"],
        [[r["id"], r["original_name"], r["sha256"]] for r in content["sources"]],
    )
    sheet(
        "نسخه‌های مبنا",
        ["حوزه", "شناسه"],
        [
            ["داده", manifest["id"]],
            ["پیش‌بینی", (content.get("forecast") or {}).get("id")],
            ["تسویه", (content.get("settlement") or {}).get("id")],
            ["گردش نقد", (content.get("cashflow") or {}).get("id")],
        ]
        + [[r["kind"], r["id"]] for r in content["reconciliation_runs"]]
        + [["تصمیم", r["id"]] for r in content["decisions"]],
    )
    sheet(
        "شاخص‌ها",
        ["شاخص", "مبنا", "مقدار دقیق " + unit_fa, "وضعیت", "علت", "مقایسه"],
        [
            [
                METRIC_LABELS.get(m["metric_key"], m["metric_key"]),
                m["basis"],
                display_amount(m["value"], unit),
                m["status"],
                ";".join(m["reason_codes"]),
                json.dumps(m.get("comparison"), ensure_ascii=False),
            ]
            for m in content["metrics"]
        ],
    )
    sheet(
        "تطبیق‌ها",
        ["شناسه", "حوزه", "نتیجه", "تصمیم", "شواهد"],
        [
            [
                c["id"],
                c["kind"],
                c["classification"],
                c.get("latest_decision_id"),
                json.dumps(c["evidence"], ensure_ascii=False),
            ]
            for c in content["candidates"]
        ],
    )
    details = []
    for row in content["rows"]:
        values = {"canonical": row["canonical_values"], "original": row["raw_values"]}
        for field, value in values.items():
            text = json.dumps(value, ensure_ascii=False)
            for part in range(0, max(1, len(text)), 30000):
                details.append(
                    [
                        row["source_file_id"],
                        row["sheet_name"],
                        row["row_number"],
                        row["source_kind"],
                        row["state"],
                        field,
                        str(part // 30000 + 1),
                        text[part : part + 30000],
                    ]
                )
    sheet(
        "جزئیات منبع",
        ["فایل", "برگه", "ردیف", "خانواده", "وضعیت", "نوع داده", "بخش", "محتوا"],
        details,
    )
    sheet(
        "شواهد",
        ["شناسه", "ردیف منبع", "قاعده", "محاسبه"],
        [
            [
                e["id"],
                e["source_row_id"],
                e["rule_key"],
                json.dumps(e["calculation"], ensure_ascii=False),
            ]
            for e in content["evidence"]
        ],
    )
    forecast = (content.get("forecast") or {}).get("content") or {}
    sheet(
        "پیش‌بینی",
        [
            "روز",
            "آغاز دقیق " + unit_fa,
            "دریافت",
            "پرداخت",
            "پایان دقیق " + unit_fa,
            "وضعیت",
            "علت",
        ],
        [
            [
                d["date"],
                display_amount(d["opening"], unit),
                display_amount(d["expected_in"], unit),
                display_amount(d["expected_out"], unit),
                display_amount(d["closing"], unit),
                d["status"],
                ";".join(d["reason_codes"]),
            ]
            for d in forecast.get("days", [])
        ],
    )
    sheet(
        "ردیف‌های خطادار",
        ["فایل", "برگه", "ردیف", "وضعیت", "علت"],
        [
            [
                r["source_file_id"],
                r["sheet_name"],
                r["row_number"],
                r["state"],
                ";".join(r["error_codes"]),
            ]
            for r in content["rows"]
            if r["state"] != "valid" or r["error_codes"]
        ],
    )
    out = BytesIO()
    book.save(out)
    return out.getvalue()
