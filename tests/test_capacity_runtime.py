"""Measured synthetic five-user upload + real scanner/typed bank financial capacity.

20,000 physical CSV rows (10 headers, 10 openings, 19,980 movements), not a
claim about every accounting layout or five simultaneous heavy analysis jobs.
"""

from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
import hashlib
import json
import os
from pathlib import Path
import statistics
import resource
import sys
from io import BytesIO
import openpyxl
import time
import unittest

import psycopg
from test_financial_runtime import FinancialTests
from test_runtime import HASHER, PASSWORD
from didban.adapters.scanner import ClamScanner
from didban.application.preparation_worker import PreparationWorker
from didban.application.analysis_worker import AnalysisWorker
from didban.application.report_worker import ReportWorker
from didban.application.financial_worker import FinancialWorker
from didban.infrastructure.financial import digest

ROOT = Path(__file__).resolve().parents[1]
MEASUREMENT = {}


class CapacityTests(FinancialTests):
    def worker(self, scanner=None):
        return PreparationWorker(
            self.worker_settings,
            scanner=scanner
            or ClamScanner("127.0.0.1", int(os.environ["DIDBAN_TEST_SCANNER_PORT"])),
        )

    def test_five_users_ten_files_twenty_thousand_physical_rows(self):
        # Two additional synthetic people, distinct from the three existing operators.
        for name in ("capacity1", "capacity2"):
            with psycopg.connect(self.dsn) as conn:
                user = conn.execute(
                    "INSERT INTO app.users(email,display_name,password_hash,activated_at) VALUES(%s,%s,%s,clock_timestamp()) RETURNING id",
                    (name + "@fixture.invalid", name, HASHER.hash(PASSWORD)),
                ).fetchone()[0]
            result = self.mutate(
                self.owner,
                "POST",
                self.base + "/memberships",
                {"user_id": str(user), "role": "advisor_accountant"},
            )
            self.assertEqual(result.status_code, 201, result.text)
        people = [
            self.client(name)
            for name in ("owner", "finance", "advisor", "capacity1", "capacity2")
        ]
        period = self.period(starts_on="2026-09-01", ends_on="2026-09-30")
        session = self.session(period)
        files = []
        for i in range(10):
            account = f"BANK{i:02d}"
            lines = [
                "row_id,date,amount,balance,reference,bank_account,currency,row_type",
                f"OPEN{i},2026-09-01,0,1000,,{account},IRR,opening",
            ]
            for j in range(1998):
                day = date(2026, 9, 1) + timedelta(days=j // 70)
                lines.append(
                    f"B{i}-{j},{day},1,{1001 + j},R{i}-{j},{account},IRR,movement"
                )
            files.append(
                ("capacity-" + str(i) + ".csv", ("\n".join(lines) + "\n").encode())
            )
        latencies, retries = [], []

        def send(pair):
            client, indexes = pair
            for i in indexes:
                started = time.monotonic()
                for attempt in range(30):
                    result = self.upload(
                        session, files[i][1], name=files[i][0], client=client
                    )
                    if result.status_code == 202:
                        latencies.append(time.monotonic() - started)
                        retries.append(attempt)
                        break
                    self.assertEqual(result.status_code, 409, result.text)
                else:
                    self.fail("Version conflict did not recover")

        start = time.monotonic()
        with ThreadPoolExecutor(max_workers=5) as executor:
            list(
                executor.map(
                    send, [(p, [i * 2, i * 2 + 1]) for i, p in enumerate(people)]
                )
            )
        uploaded_seconds = time.monotonic() - start
        self.assertEqual(len(self.files(session)), 10)
        self.assertEqual(self.seal(session).status_code, 202)
        interpreted_start = time.monotonic()
        for _ in range(120):
            state = self.detail(session)
            if state["interpretation_finished_at"]:
                break
            if state["status"] == "waiting_input":
                issue = self.issue(session)
                self.assertEqual(issue["code"], "interpretation_coverage", issue)
                self.assertEqual(self.answer(session, "unknown").status_code, 201)
            if not self.worker().run_once():
                self.interpreter().run_once()
        else:
            self.fail("Capacity interpretation did not finish")
        files_view = self.files(session)
        self.assertEqual(
            sum(f["preparation_summary"]["nonempty_rows"] for f in files_view), 20000
        )
        self.assertEqual(
            sum(f["interpretation_summary"]["valid_rows"] for f in files_view), 19990
        )
        self.assertTrue(
            all(
                "ClamAV" in f.get("scan_engine", "") or f["status"] == "ready"
                for f in files_view
            )
        )
        with psycopg.connect(self.dsn) as conn:
            bindings = {
                "binding|" + digest(str(file_id) + ":" + sheet)[:20]: account
                for file_id, sheet, account in conn.execute(
                    "SELECT DISTINCT i.source_file_id,i.sheet_name,n.canonical_values->>'bank_account' FROM app.normalized_source_rows n JOIN app.source_rows s ON s.id=n.source_row_id AND s.company_id=n.company_id JOIN app.interpretation_versions i ON i.id=s.interpretation_id AND i.company_id=s.company_id WHERE n.company_id=%s AND n.source_kind='bank'",
                    (self.company,),
                )
            }
        financial_start = time.monotonic()
        self.start_financial(session)
        worker = FinancialWorker(self.worker_settings)
        questions = 0
        for _ in range(120):
            state = self.detail(session)
            if state.get("manifest_id"):
                break
            if state["status"] == "waiting_input":
                issue = self.issue(session)
                choices = issue["constraints"]["choices"]
                key = issue["constraints"].get("financial_key", "")
                if key.startswith("binding|"):
                    value = bindings[key]
                else:
                    value = next(
                        (
                            c
                            for c in (
                                "declared_complete",
                                "affirmed",
                                "start_of_day",
                                "zero_confirmed",
                                "not_required",
                            )
                            if c in choices
                        ),
                        choices[0],
                    )
                self.assertEqual(self.answer(session, value).status_code, 201)
                questions += 1
            self.assertTrue(worker.run_once())
        else:
            self.fail("Capacity financial result did not publish")
        metrics = self.metrics(state["manifest_id"])
        self.assertEqual(metrics["cash"]["value"], "29980")
        self.assertIsNone(metrics["net_profit"]["value"])
        financial_seconds = time.monotonic() - financial_start
        MEASUREMENT.update(
            {
                "completed_stage": "core_financial",
                "attempted_stage": "business_analysis",
                "distinct_users": 5,
                "simultaneous_upload_clients": 5,
                "file_count": 10,
                "physical_rows": 20000,
                "typed_bank_rows": 19990,
                "movements": 19980,
                "exact_cash_irr": "29980",
                "upload_seconds": round(uploaded_seconds, 3),
                "interpretation_seconds": round(financial_start - interpreted_start, 3),
                "financial_seconds": round(financial_seconds, 3),
                "model_calls": 0,
                "scope_limit": "one shared synthetic bank session; explicit stage starts in integration harness; concurrent uploads, not five concurrent heavy analyses; missing ledger profit remains null",
            }
        )
        reporting_start = time.monotonic()
        started_analysis = self.mutate(
            self.owner,
            "POST",
            self.base + "/processing-sessions/" + session["id"] + "/analyze",
            None,
            version=self.detail(session)["session_version"],
        )
        self.assertEqual(started_analysis.status_code, 202, started_analysis.text)
        analysis = AnalysisWorker(self.worker_settings)
        self.assertTrue(analysis.run_once())
        analysis_state = self.detail(session)
        self.assertIsNotNone(analysis_state["analysis_finished_at"])
        self.assertNotEqual(analysis_state["status"], "waiting_input")
        MEASUREMENT.update(
            {
                "completed_stage": "business_analysis",
                "attempted_stage": "report_generation",
                "analysis_seconds": round(time.monotonic() - reporting_start, 3),
            }
        )
        renderer = ReportWorker(self.worker_settings)
        for _ in range(3):
            self.assertTrue(renderer.run_once())
        reports = self.owner.get(self.base + "/reports?limit=100")
        self.assertEqual(reports.status_code, 200, reports.text)
        edition = next(
            r
            for r in reports.json()["items"]
            if r["manifest_id"] == state["manifest_id"]
        )
        self.assertEqual(
            {e["format"] for e in edition["exports"]}, {"pdf", "xlsx", "error_csv"}
        )
        artifact_sizes = {}
        for artifact in edition["exports"]:
            self.assertEqual(artifact["status"], "ready", artifact)
            downloaded = self.owner.get(
                self.base + "/exports/" + artifact["id"] + "/content"
            )
            self.assertEqual(
                downloaded.status_code,
                200,
                downloaded.text[:100] if downloaded.status_code != 200 else "",
            )
            artifact_sizes[artifact["format"]] = len(downloaded.content)
            self.assertEqual(
                hashlib.sha256(downloaded.content).hexdigest(), artifact["sha256"]
            )
            if artifact["format"] == "xlsx":
                book = openpyxl.load_workbook(
                    BytesIO(downloaded.content), read_only=True, data_only=False
                )
                self.assertGreaterEqual(book["جزئیات منبع"].max_row, 19991)
                metric_rows = book["شاخص‌ها"].iter_rows(values_only=True)
                headers = next(metric_rows)
                amount_index = headers.index("مقدار دقیق ریال")
                self.assertIn(
                    "29980",
                    [row[amount_index] for row in metric_rows],
                )
                book.close()
        self.assertEqual(self.detail(session)["status"], "completed_limited")

        def read_result(client):
            started = time.monotonic()
            response = client.get(
                self.base + "/results/" + state["manifest_id"] + "/metrics?limit=100"
            )
            self.assertEqual(response.status_code, 200)
            return time.monotonic() - started

        with ThreadPoolExecutor(max_workers=5) as executor:
            read_latencies = list(executor.map(read_result, people))
        peak_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        peak_rss_mib = peak_rss / (1024**2 if sys.platform == "darwin" else 1024)
        MEASUREMENT.update(
            {
                "status": "passed",
                "completed_stage": "report_generation_and_download",
                "distinct_users": 5,
                "simultaneous_upload_clients": 5,
                "file_count": 10,
                "physical_rows": 20000,
                "typed_bank_rows": 19990,
                "movements": 19980,
                "exact_cash_irr": "29980",
                "upload_seconds": round(uploaded_seconds, 3),
                "interpretation_seconds": round(financial_start - interpreted_start, 3),
                "financial_seconds": round(financial_seconds, 3),
                "report_generation_and_download_seconds": round(
                    time.monotonic() - reporting_start, 3
                ),
                "artifact_bytes": artifact_sizes,
                "test_process_peak_rss_mib": round(peak_rss_mib, 1),
                "five_user_result_read_max_seconds": round(max(read_latencies), 3),
                "report_exports_ready": 3,
                "upload_response_median_seconds": round(
                    statistics.median(latencies), 3
                ),
                "upload_version_conflict_retries": sum(retries),
                "financial_questions": questions,
                "scanner": "real ClamAV",
                "model_calls": 0,
                "scope_limit": "one shared synthetic bank session; concurrent uploads, not five concurrent heavy analyses; missing ledger profit remains null",
            }
        )


if __name__ == "__main__":
    suite = unittest.TestSuite(
        [CapacityTests("test_five_users_ten_files_twenty_thousand_physical_rows")]
    )
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    MEASUREMENT["status"] = "passed" if result.wasSuccessful() else "failed"
    peak_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    MEASUREMENT["test_process_peak_rss_mib"] = round(
        peak_rss / (1024**2 if sys.platform == "darwin" else 1024), 1
    )
    MEASUREMENT["source_sha256"] = {
        str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in [Path(__file__).resolve()]
    }
    (ROOT / "docs/capacity-validation.json").write_text(
        json.dumps(MEASUREMENT, indent=2) + "\n"
    )
    raise SystemExit(not result.wasSuccessful())
