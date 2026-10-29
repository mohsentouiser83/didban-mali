"""Real company isolation, immutable report editions and private managerial exports."""

import base64
from io import BytesIO
import json
from pathlib import Path
import unittest
from openpyxl import load_workbook
from test_reports_runtime import ReportTests


class ManagementRuntimeTests(ReportTests):
    def test_management_snapshot_and_exports_are_pinned(self):
        session, report = self.report_fixture()
        url = self.base + "/reports/" + report["id"] + "/management"
        result = self.owner.get(url)
        self.assertEqual(result.status_code, 200, result.text)
        body = result.json()
        self.assertEqual(
            next(t for t in body["statements"]["totals"] if t["key"] == "net_profit")[
                "value"
            ],
            "448",
        )
        self.assertIsNotNone(body["analysis"]["profit_delta"])
        self.assertIsNotNone(body["analysis"]["cash"]["change"]["value"])
        self.assertEqual(self.client("viewer").get(url).status_code, 200)
        self.assertEqual(self.client("outsider").get(url).status_code, 404)
        self.drain()
        exact = self.owner.get(self.base + "/reports/" + report["id"]).json()
        job = next(j for j in exact["exports"] if j["format"] == "xlsx")
        download = self.owner.get(self.base + "/exports/" + job["id"] + "/content")
        book = load_workbook(BytesIO(download.content))
        self.assertIn("صورت‌های مالی", book.sheetnames)
        self.assertIn("گزارش روزانه", book.sheetnames)
        self.assertFalse(
            any(c.data_type == "f" for ws in book for row in ws for c in row)
        )
        # Shortened report is a new immutable edition; the old edition never changes.
        revised = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"as_of_date": "2026-09-15"},
        )
        self.assertEqual(revised.status_code, 201, revised.text)
        self.assertNotEqual(revised.json()["id"], report["id"])
        self.assertEqual(revised.json()["as_of_date"], "2026-09-15")
        self.assertIsNone(
            next(
                metric
                for metric in revised.json()["metrics"]
                if metric["metric_key"] == "net_profit"
            )["value"]
        )
        self.assertIsNone(revised.json()["forecast_closing"])
        self.assertEqual(revised.json()["approval_ids"], [])
        shortened = self.owner.get(
            self.base + "/reports/" + revised.json()["id"] + "/management"
        ).json()
        self.assertEqual(shortened["as_of_date"], "2026-09-15")
        self.assertEqual(len(shortened["daily"]["days"]), 15)
        shifted = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"starts_on": "2026-09-10", "as_of_date": "2026-09-15"},
        )
        self.assertEqual(shifted.status_code, 201, shifted.text)
        shifted_content = self.owner.get(
            self.base + "/reports/" + shifted.json()["id"] + "/management"
        ).json()
        self.assertEqual(
            shifted_content["analysis"]["comparison_starts_on"], "2026-08-10"
        )
        self.assertEqual(
            shifted_content["analysis"]["comparison_as_of_date"], "2026-08-15"
        )
        self.assertEqual(self.owner.get(url).json(), body)
        self.drain()
        revised_detail = self.owner.get(
            self.base + "/reports/" + revised.json()["id"]
        ).json()
        xlsx_job = next(
            job for job in revised_detail["exports"] if job["format"] == "xlsx"
        )
        revised_book = load_workbook(
            BytesIO(
                self.owner.get(
                    self.base + "/exports/" + xlsx_job["id"] + "/content"
                ).content
            )
        )
        self.assertIn("شاخص‌های نسخه داده", revised_book.sheetnames)
        headers = [cell.value for row in revised_book["خلاصه"] for cell in row]
        self.assertIn("2026-09-15", headers)
        self.assertIn("شواهد شاخص‌های بازه", revised_book.sheetnames)
        self.assertEqual(self.detail(session)["status"], "completed_limited")

    def test_custom_report_does_not_replace_active_session_report(self):
        session, report = self.report_fixture()
        revised = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"as_of_date": "2026-09-15"},
        )
        self.assertEqual(revised.status_code, 201, revised.text)
        self.assertEqual(self.detail(session)["report_id"], report["id"])
        self.drain()
        self.assertEqual(self.detail(session)["report_id"], report["id"])
        self.assertEqual(self.detail(session)["status"], "completed_limited")

    def test_template_bounds_permissions_and_replay(self):
        _, report = self.report_fixture()
        with self.app.state.db.transaction(
            identity=False, actor=self.users["owner"], company=self.company
        ) as conn:
            constraint = conn.execute(
                "SELECT pg_get_constraintdef(oid) AS definition FROM pg_constraint WHERE conrelid='app.account_role_versions'::regclass AND conname='account_role_versions_role_check'"
            ).fetchone()["definition"]
            for role in (
                "fixed_asset",
                "other_asset",
                "contra_asset",
                "other_equity",
                "bank",
                "capital",
            ):
                self.assertIn(role, constraint)
        filled = self.owner.get(
            self.base + "/reports/" + report["id"] + "/statement-template"
        )
        self.assertEqual(filled.status_code, 200, filled.text[:100])
        filled_parse = self.mutate(
            self.owner,
            "POST",
            self.base + "/statement-templates/parse",
            {
                "name": "قالب گزارش",
                "content_base64": base64.b64encode(filled.content).decode(),
            },
        )
        self.assertEqual(filled_parse.status_code, 201, filled_parse.text)
        self.assertTrue(any(line["accounts"] for line in filled_parse.json()["lines"]))
        complete_template = filled_parse.json()
        templated = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"template": complete_template},
        )
        self.assertEqual(templated.status_code, 201, templated.text)
        templated_content = self.owner.get(
            self.base + "/reports/" + templated.json()["id"] + "/management"
        ).json()
        self.assertEqual(templated_content["statements"]["unmapped_accounts"], [])
        self.assertEqual(
            next(
                t
                for t in templated_content["statements"]["totals"]
                if t["key"] == "net_profit"
            )["value"],
            "448",
        )
        self.assertIsNotNone(templated_content["analysis"]["profit_delta"])
        example = self.owner.get(self.base + "/statement-templates/example")
        self.assertEqual(example.status_code, 200, example.text[:100])
        payload = {
            "name": "قالب مرجع",
            "content_base64": base64.b64encode(example.content).decode(),
        }
        parsed = self.mutate(
            self.owner, "POST", self.base + "/statement-templates/parse", payload
        )
        self.assertEqual(parsed.status_code, 201, parsed.text)
        self.assertEqual(len(parsed.json()["lines"]), 8)
        denied = self.mutate(
            self.client("viewer"),
            "POST",
            self.base + "/statement-templates/parse",
            payload,
        )
        self.assertEqual(denied.status_code, 403)
        template = parsed.json()
        template["lines"][0]["accounts"] = ["حساب ناموجود"]
        path = self.base + "/results/" + report["manifest_id"] + "/management-reports"
        key = "management-idempotency-test"
        created = self.mutate(self.owner, "POST", path, {"template": template}, key=key)
        self.assertEqual(created.status_code, 201, created.text)
        replay = self.mutate(self.owner, "POST", path, {"template": template}, key=key)
        self.assertEqual(replay.json()["id"], created.json()["id"])
        payload = self.owner.get(
            self.base + "/reports/" + created.json()["id"] + "/management"
        ).json()
        self.assertIsNone(payload["statements"]["lines"][0]["value"])
        for body in (
            {"starts_on": "2026-08-01"},
            {"as_of_date": "2026-10-01"},
            {"starts_on": "2026-09-20", "as_of_date": "2026-09-10"},
        ):
            rejected = self.mutate(self.owner, "POST", path, body)
            self.assertEqual(rejected.status_code, 422, rejected.text)
        denied = self.mutate(self.client("viewer"), "POST", path, {})
        self.assertEqual(denied.status_code, 403)


if __name__ == "__main__":
    suite = unittest.TestSuite(
        ManagementRuntimeTests(name)
        for name in ManagementRuntimeTests.__dict__
        if name.startswith("test_")
    )
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    Path("docs/management-runtime-validation.json").write_text(
        json.dumps(
            {
                "status": "passed"
                if result.wasSuccessful() and not result.skipped
                else "failed_or_not_run",
                "tests_run": result.testsRun,
                "errors": len(result.errors),
                "failures": len(result.failures),
                "skipped": len(result.skipped),
                "live_model_calls": 0,
            },
            indent=2,
        )
        + "\n"
    )
    raise SystemExit(0 if result.wasSuccessful() and not result.skipped else 1)
