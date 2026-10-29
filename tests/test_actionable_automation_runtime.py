"""Automatic immutable summaries, custom dates and bounded deadline observations."""

from io import BytesIO
import unittest
import psycopg
from openpyxl import load_workbook
from test_reports_runtime import ReportTests


class ActionableAutomationTests(ReportTests):
    def test_report_publishes_generic_summary_and_exports_automatically(self):
        _, report = self.report_fixture()
        result = self.client("viewer").get(
            self.base + "/reports/" + report["id"] + "/actionable-analysis"
        )
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(
            {p["axis"] for p in result.json()["items"]},
            {"cash", "balances", "profit", "quality"},
        )
        latest = self.owner.get(self.base + "/actionable-summary")
        self.assertEqual(latest.status_code, 200, latest.text)
        self.assertEqual(latest.json()["report_id"], report["id"])
        with psycopg.connect(self.dsn) as conn:
            content = conn.execute(
                "SELECT content FROM app.report_snapshots WHERE id=%s", (report["id"],)
            ).fetchone()[0]
        self.assertNotIn(
            "payroll", {p["axis"] for p in content["actionable_summary"]["items"]}
        )
        self.drain()
        exact = self.owner.get(self.base + "/reports/" + report["id"]).json()
        job = next(j for j in exact["exports"] if j["format"] == "xlsx")
        book = load_workbook(
            BytesIO(
                self.owner.get(self.base + "/exports/" + job["id"] + "/content").content
            )
        )
        self.assertIn("پیشنهادهای مدیریتی", book.sheetnames)

    def test_short_report_keeps_actual_scope_embedded_values_and_source_rows(self):
        _, original = self.report_fixture()
        result = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + original["manifest_id"] + "/management-reports",
            {"as_of_date": "2026-09-15"},
        )
        self.assertEqual(result.status_code, 201, result.text)
        report = result.json()
        with psycopg.connect(self.dsn) as conn:
            content = conn.execute(
                "SELECT content FROM app.report_snapshots WHERE id=%s", (report["id"],)
            ).fetchone()[0]
        proposals = content["actionable_summary"]["items"]
        self.assertTrue(
            all(p["scope"]["as_of_date"] == "2026-09-15" for p in proposals)
        )
        self.assertTrue(all(not p["due_evidence"] for p in proposals))
        metric = next(m for m in content["metrics"] if m["metric_key"] == "revenue")
        profit = next(p for p in proposals if p["axis"] == "profit")
        evidence = next(e for e in profit["evidence"] if e["metric_key"] == "revenue")
        self.assertEqual(evidence["value"], metric["value"])
        self.assertTrue(evidence["source_refs"])
        self.assertTrue(
            all(
                ref["file_name"] and ref["row_number"]
                for ref in evidence["source_refs"]
            )
        )
        self.assertTrue(any(ref["input_values"] for ref in evidence["source_refs"]))
        self.assertTrue(
            all(
                ref["calculated_value"] == metric["value"]
                for ref in evidence["source_refs"]
            )
        )
        self.assertEqual(content["manifest"]["as_of_date"], "2026-09-30")

    def test_future_deadline_approaches_once_and_paid_invoice_is_suppressed(self):
        _, report = self.report_fixture()
        with psycopg.connect(self.dsn) as conn:
            proposal = conn.execute(
                "SELECT id,content FROM app.actionable_proposals WHERE company_id=%s AND report_id=%s AND axis='balances'",
                (self.company, report["id"]),
            ).fetchone()
            self.assertTrue(
                any(e["due_on"] == "2026-10-10" for e in proposal[1]["due_evidence"])
            )
            conn.execute(
                "DELETE FROM app.actionable_notification_states WHERE company_id=%s",
                (self.company,),
            )
            first = conn.execute(
                "SELECT app.notify_actionable_at(%s,'2026-09-30')", (proposal[0],)
            ).fetchone()[0]
            self.assertGreater(first, 0)
            same = conn.execute(
                "SELECT app.notify_actionable_at(%s,'2026-09-30')", (proposal[0],)
            ).fetchone()[0]
            self.assertEqual(same, 0)
            approaching = conn.execute(
                "SELECT app.notify_actionable_at(%s,'2026-10-03')", (proposal[0],)
            ).fetchone()[0]
            self.assertGreater(approaching, 0)
            unchanged = conn.execute(
                "SELECT app.notify_actionable_at(%s,'2026-10-04')", (proposal[0],)
            ).fetchone()[0]
            self.assertEqual(unchanged, 0)
            # Trusted test fixture introduces a validated post-snapshot settlement;
            # original report/proposal remain untouched and immutable.
            invoice = conn.execute(
                "SELECT id,party_id,record_revision_id,gross FROM app.invoices WHERE company_id=%s AND source_invoice_id='P001'",
                (self.company,),
            ).fetchone()
            paid = conn.execute(
                "SELECT coalesce(sum(a.amount),0) FROM app.allocations a WHERE a.company_id=%s AND a.invoice_id=%s AND NOT EXISTS(SELECT 1 FROM app.allocation_reversals x WHERE x.company_id=a.company_id AND x.allocation_id=a.id)",
                (self.company, invoice[0]),
            ).fetchone()[0]
            remaining = invoice[3] - paid
            conn.execute("SET LOCAL search_path=app,pg_catalog")
            conn.execute(
                "SELECT set_config('app.user_id',%s,true),set_config('app.company_id',%s,true)",
                (self.users["owner"], self.company),
            )
            event = conn.execute(
                "INSERT INTO app.settlement_events(company_id,record_revision_id,party_id,settled_on,direction,amount,event_role) VALUES(%s,%s,%s,'2026-10-04','payment',%s,'financial') RETURNING id",
                (self.company, invoice[2], invoice[1], remaining),
            ).fetchone()[0]
            dataset = conn.execute(
                "SELECT dataset_revision_id FROM app.result_manifests WHERE id=%s",
                (report["manifest_id"],),
            ).fetchone()[0]
            conn.execute(
                "INSERT INTO app.allocations(company_id,settlement_event_id,invoice_id,dataset_revision_id,amount,actor_id,provenance,effect_key) VALUES(%s,%s,%s,%s,%s,%s,'human','deadline-test-payment')",
                (
                    self.company,
                    event,
                    invoice[0],
                    dataset,
                    remaining,
                    self.users["owner"],
                ),
            )
            suppressed = conn.execute(
                "SELECT app.notify_actionable_at(%s,'2026-10-05')", (proposal[0],)
            ).fetchone()[0]
            self.assertEqual(suppressed, 0)

    def test_disabled_module_and_superseding_report_prevent_old_reminders(self):
        _, report = self.report_fixture()
        self.drain()
        with psycopg.connect(self.dsn) as conn:
            proposal = conn.execute(
                "SELECT id FROM app.actionable_proposals WHERE company_id=%s AND report_id=%s AND axis='balances'",
                (self.company, report["id"]),
            ).fetchone()[0]
            conn.execute(
                "DELETE FROM app.actionable_notification_states WHERE company_id=%s",
                (self.company,),
            )
        modules = self.owner.get(self.base + "/modules").json()
        modules["modules"]["balances"] = False
        changed = self.mutate(
            self.owner,
            "PUT",
            self.base + "/modules",
            {"modules": modules["modules"], "reason": "آزمون توقف اعلان"},
            version=modules["version"],
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT app.notify_actionable_at(%s,'2026-10-03')", (proposal,)
                ).fetchone()[0],
                0,
            )
        modules = self.owner.get(self.base + "/modules").json()
        modules["modules"]["balances"] = True
        self.assertEqual(
            self.mutate(
                self.owner,
                "PUT",
                self.base + "/modules",
                {"modules": modules["modules"], "reason": "بازگرداندن"},
                version=modules["version"],
            ).status_code,
            200,
        )
        revised = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"as_of_date": "2026-09-15"},
        )
        self.assertEqual(revised.status_code, 201, revised.text)
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT app.notify_actionable_at(%s,'2026-10-03')", (proposal,)
                ).fetchone()[0],
                0,
            )


if __name__ == "__main__":
    suite = unittest.TestSuite(
        ActionableAutomationTests(n)
        for n in ActionableAutomationTests.__dict__
        if n.startswith("test_")
    )
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(not result.wasSuccessful())
