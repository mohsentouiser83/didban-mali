"""Restricted PostgreSQL projection, allocation and snapshot/plan regression tests."""

from pathlib import Path
import hashlib
import json
import unittest
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from concurrent.futures import ThreadPoolExecutor
import psycopg
from psycopg.rows import dict_row
from test_financial_runtime import FinancialTests
from didban.infrastructure.settlements import project

ROOT = Path(__file__).resolve().parents[1]


class SettlementTests(FinancialTests):
    def test_missing_due_date_keeps_known_balance_but_unknown_aging(self):
        original_upload = self.upload

        def upload_with_missing_due(session, data, **kwargs):
            if kwargs.get("name") == "invoices.xlsx":
                book = load_workbook(BytesIO(data))
                book["فروش"].cell(3, 4).value = None
                output = BytesIO()
                book.save(output)
                data = output.getvalue()
            return original_upload(session, data, **kwargs)

        with patch.object(self, "upload", side_effect=upload_with_missing_due):
            manifest = self.reference()
        work = self.owner.get(self.base + "/settlements/" + manifest["id"]).json()
        invoice = next(i for i in work["invoices"] if i["source_invoice_id"] == "S001")
        self.assertEqual(invoice["remaining_as_of"], "600")
        self.assertIsNone(invoice["due_on"])
        self.assertIsNone(invoice["overdue_days"])
        self.assertIn("due_date_missing", invoice["reason_codes"])

    def test_aging_recomputed_after_reversal_and_full_settlement(self):
        manifest = self.reference()
        path = self.base + "/settlements/" + manifest["id"]
        work = self.owner.get(path).json()
        invoice = next(i for i in work["invoices"] if i["source_invoice_id"] == "S000")
        self.assertEqual(invoice["remaining_as_of"], "0")
        self.assertEqual(invoice["overdue_days"], 0)
        allocation = next(
            a for a in work["allocations"] if a["invoice_id"] == invoice["invoice_id"]
        )
        reversed_result = self.mutate(
            self.owner,
            "POST",
            self.base + "/allocations/" + allocation["id"] + "/reversals",
            {"reason": "آزمون سررسید پس از برگشت"},
            version=allocation["version"],
        )
        self.assertEqual(reversed_result.status_code, 201, reversed_result.text)
        after = self.owner.get(path).json()
        restored_balance = next(
            i for i in after["invoices"] if i["invoice_id"] == invoice["invoice_id"]
        )
        self.assertEqual(restored_balance["remaining_as_of"], "300")
        self.assertEqual(restored_balance["overdue_days"], 51)
        result = self.mutate(
            self.owner,
            "POST",
            self.base + "/allocations",
            {
                "invoice_id": invoice["invoice_id"],
                "settlement_event_id": allocation["settlement_event_id"],
                "dataset_revision_id": manifest["dataset_revision_id"],
                "amount": "300",
                "reason": "تسویه دوباره برای آزمون",
            },
        )
        self.assertEqual(result.status_code, 201, result.text)
        final = next(
            i
            for i in self.owner.get(path).json()["invoices"]
            if i["invoice_id"] == invoice["invoice_id"]
        )
        self.assertEqual(final["remaining_as_of"], "0")
        self.assertEqual(final["overdue_days"], 0)
        self.assertEqual(
            next(
                i for i in work["invoices"] if i["invoice_id"] == invoice["invoice_id"]
            )["overdue_days"],
            0,
        )

    def reference(self):
        FinancialTests.test_reference_baseline_exact_metrics_and_comparison(self)
        manifest = self.owner.get(self.base + "/results").json()["items"][0]
        with psycopg.connect(self.worker_settings.dsn, row_factory=dict_row) as conn:
            conn.execute(
                "SELECT set_config('search_path','app,pg_catalog',true),set_config('app.user_id',%s,true),set_config('app.company_id',%s,true)",
                (self.users["owner"], self.company),
            )
            project(conn, self.company, manifest["id"])
        return manifest

    def test_imported_events_once_and_reversal_creates_new_workspace_only(self):
        manifest = self.reference()
        path = self.base + "/settlements/" + manifest["id"]
        first = self.owner.get(path)
        self.assertEqual(first.status_code, 200, first.text)
        work = first.json()
        self.assertEqual(len(work["allocations"]), 4)
        old_invoices = self.owner.get(
            self.base + "/results/" + manifest["id"] + "/receivables"
        ).json()
        invoice = next(i for i in work["invoices"] if i["source_invoice_id"] == "S000")
        allocation = next(
            a for a in work["allocations"] if a["invoice_id"] == invoice["invoice_id"]
        )
        result = self.mutate(
            self.client("advisor"),
            "POST",
            self.base + "/allocations/" + allocation["id"] + "/reversals",
            {"reason": "آزمون برگشت ارتباط"},
            version=allocation["version"],
        )
        self.assertEqual(result.status_code, 201, result.text)
        after = self.owner.get(path).json()
        self.assertNotEqual(after["id"], work["id"])
        self.assertEqual(
            next(i for i in after["invoices"] if i["source_invoice_id"] == "S000")[
                "remaining_as_of"
            ],
            "300",
        )
        self.assertEqual(
            self.owner.get(
                self.base + "/results/" + manifest["id"] + "/receivables"
            ).json(),
            old_invoices,
        )
        with psycopg.connect(self.worker_settings.dsn, row_factory=dict_row) as conn:
            conn.execute(
                "SELECT set_config('search_path','app,pg_catalog',true),set_config('app.user_id',%s,true),set_config('app.company_id',%s,true)",
                (self.users["owner"], self.company),
            )
            project(conn, self.company, manifest["id"])
        self.assertEqual(self.owner.get(path).json()["version"], after["version"])

    def test_capacity_replay_direction_and_viewer_are_enforced(self):
        manifest = self.reference()
        work = self.owner.get(self.base + "/settlements/" + manifest["id"]).json()
        inv = next(i for i in work["invoices"] if i["source_invoice_id"] == "S001")
        event = next(
            e
            for e in work["events"]
            if e["direction"] == "receipt" and e["allocated"] == "400"
        )
        body = {
            "invoice_id": inv["invoice_id"],
            "settlement_event_id": event["id"],
            "dataset_revision_id": manifest["dataset_revision_id"],
            "amount": "1",
            "reason": "بیش از ظرفیت",
        }
        self.assertEqual(
            self.mutate(
                self.owner, "POST", self.base + "/allocations", body
            ).status_code,
            409,
        )
        self.assertEqual(
            self.mutate(
                self.client("viewer"), "POST", self.base + "/allocations", body
            ).status_code,
            403,
        )
        allocation = next(
            a for a in work["allocations"] if a["settlement_event_id"] == event["id"]
        )
        self.assertEqual(
            self.mutate(
                self.owner,
                "POST",
                self.base + "/allocations/" + allocation["id"] + "/reversals",
                {"reason": "آزمون"},
                version=1,
            ).status_code,
            201,
        )
        body["amount"] = "400"
        body["reason"] = "ارتباط معتبر"
        key = "fixture-allocation-replay"
        one = self.mutate(self.owner, "POST", self.base + "/allocations", body, key=key)
        two = self.mutate(self.owner, "POST", self.base + "/allocations", body, key=key)
        self.assertEqual(one.status_code, 201, one.text)
        self.assertEqual(one.json(), two.json())
        other = self.client("outsider")
        self.assertEqual(
            other.get(self.base + "/settlements/" + manifest["id"]).status_code, 404
        )

    def test_concurrent_allocations_cannot_overspend_the_same_event(self):
        manifest = self.reference()
        work = self.owner.get(self.base + "/settlements/" + manifest["id"]).json()
        invoice = next(i for i in work["invoices"] if i["source_invoice_id"] == "S001")
        event = next(e for e in work["events"] if e["allocated"] == "400")
        allocation = next(
            a for a in work["allocations"] if a["settlement_event_id"] == event["id"]
        )
        self.assertEqual(
            self.mutate(
                self.owner,
                "POST",
                self.base + "/allocations/" + allocation["id"] + "/reversals",
                {"reason": "آزمون ظرفیت هم‌زمان"},
                version=1,
            ).status_code,
            201,
        )
        body = {
            "invoice_id": invoice["invoice_id"],
            "settlement_event_id": event["id"],
            "dataset_revision_id": manifest["dataset_revision_id"],
            "amount": "300",
            "reason": "آزمون ظرفیت هم‌زمان",
        }
        clients = [self.client("owner"), self.client("owner")]
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(
                pool.map(
                    lambda client: self.mutate(
                        client, "POST", self.base + "/allocations", body
                    ),
                    clients,
                )
            )
        self.assertEqual(
            sorted(r.status_code for r in results),
            [201, 409],
            [r.text for r in results],
        )

    def test_same_file_reinterpretation_preserves_manual_effect_and_capacity(self):
        manifest = self.reference()
        work = self.owner.get(self.base + "/settlements/" + manifest["id"]).json()
        invoice = next(i for i in work["invoices"] if i["source_invoice_id"] == "S001")
        event = next(e for e in work["events"] if e["allocated"] == "400")
        allocation = next(
            a for a in work["allocations"] if a["settlement_event_id"] == event["id"]
        )
        self.assertEqual(
            self.mutate(
                self.owner,
                "POST",
                self.base + "/allocations/" + allocation["id"] + "/reversals",
                {"reason": "اصلاح ارتباط"},
                version=1,
            ).status_code,
            201,
        )
        manual = self.mutate(
            self.owner,
            "POST",
            self.base + "/allocations",
            {
                "invoice_id": invoice["invoice_id"],
                "settlement_event_id": event["id"],
                "dataset_revision_id": manifest["dataset_revision_id"],
                "amount": "400",
                "reason": "ارتباط انسانی معتبر",
            },
        )
        self.assertEqual(manual.status_code, 201, manual.text)
        session = self.session({"id": manifest["period_id"]})
        for name in ("ledger.xlsx", "bank.xlsx", "invoices.xlsx", "settlements.csv"):
            path = ROOT / "docs/validation/rebuild-v1/inputs/baseline" / name
            self.assertEqual(
                self.upload(session, path.read_bytes(), name=name).status_code, 202
            )
        self.assertEqual(self.seal(session).status_code, 202)
        self.complete_coverage(session)
        self.start_financial(session)
        from didban.application.financial_worker import FinancialWorker

        worker = FinancialWorker(self.worker_settings)
        for _ in range(100):
            state = self.detail(session)
            if state["manifest_id"]:
                break
            if state["status"] == "waiting_input":
                issue = self.issue(session)
                key = issue["constraints"]["financial_key"]
                choices = issue["constraints"]["choices"]
                if key.startswith("income_tax|"):
                    value = (
                        "recognized"
                        if key.endswith(manifest["period_id"])
                        else "zero_confirmed"
                    )
                elif key.startswith("cogs|"):
                    value = (
                        "recognized"
                        if key.endswith(manifest["period_id"])
                        else "not_required"
                    )
                elif key.startswith("binding|"):
                    value = (
                        "بانک الف" if "بانک الف" in issue["question_fa"] else "بانک ب"
                    )
                else:
                    value = next(
                        c
                        for c in (
                            "declared_complete",
                            "affirmed",
                            "recognized",
                            "start_of_day",
                        )
                        if c in choices
                    )
                self.assertEqual(self.answer(session, value).status_code, 201)
            self.assertTrue(worker.run_once())
        else:
            self.fail("second financial revision did not finish")
        with psycopg.connect(self.worker_settings.dsn, row_factory=dict_row) as conn:
            conn.execute(
                "SELECT set_config('search_path','app,pg_catalog',true),set_config('app.user_id',%s,true),set_config('app.company_id',%s,true)",
                (self.users["owner"], self.company),
            )
            project(conn, self.company, state["manifest_id"])
        current = self.owner.get(
            self.base + "/settlements/" + state["manifest_id"]
        ).json()
        self.assertIn(manual.json()["id"], [a["id"] for a in current["allocations"]])
        self.assertEqual(
            next(i for i in current["invoices"] if i["source_invoice_id"] == "S001")[
                "remaining_as_of"
            ],
            "600",
        )
        used = next(e for e in current["events"] if e["id"] == event["id"])
        self.assertEqual(used["allocated"], "400")

    def test_plans_are_exact_versioned_and_old_revision_is_not_overwritten(self):
        manifest = self.reference()
        body = {
            "title": "اجاره",
            "scheduled_on": "2026-10-20",
            "total_amount": "12.345678",
            "unit": "IRT",
        }
        first = self.mutate(
            self.client("advisor"), "POST", self.base + "/planned-payments", body
        )
        self.assertEqual(first.status_code, 201, first.text)
        p = first.json()
        self.assertEqual(p["total_amount"], "123.45678")
        self.assertEqual(p["as_of_date"], manifest["as_of_date"])
        work = self.owner.get(self.base + "/settlements/" + manifest["id"]).json()
        invoice = next(i for i in work["invoices"] if i["direction"] == "purchase")
        linked = self.mutate(
            self.owner,
            "POST",
            self.base + "/planned-payments",
            body | {"invoice_id": invoice["invoice_id"]},
        )
        self.assertEqual(linked.status_code, 201, linked.text)
        # Linked and independent plans share a list. Pagination uses plan IDs,
        # never nullable/shared invoice IDs, and must not skip either record.
        page = self.owner.get(self.base + "/planned-payments?limit=1")
        self.assertEqual(page.status_code, 200, page.text)
        following = self.owner.get(
            self.base + "/planned-payments?limit=1&cursor=" + page.json()["next_cursor"]
        )
        self.assertEqual(following.status_code, 200, following.text)
        self.assertEqual(
            {page.json()["items"][0]["id"], following.json()["items"][0]["id"]},
            {p["id"], linked.json()["id"]},
        )
        unit_only = self.mutate(
            self.owner,
            "PATCH",
            self.base + "/planned-payments/" + p["id"],
            {"unit": "IRT"},
            version=p["version"],
        )
        self.assertEqual(unit_only.status_code, 422, unit_only.text)
        edited = self.mutate(
            self.owner,
            "PATCH",
            self.base + "/planned-payments/" + p["id"],
            {"title": "اجاره مهر", "total_amount": "200", "unit": "IRR"},
            version=p["version"],
        )
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertNotEqual(edited.json()["id"], p["id"])
        stale = self.mutate(
            self.owner,
            "PATCH",
            self.base + "/planned-payments/" + p["id"],
            {"title": "قدیمی"},
            version=p["version"],
        )
        self.assertEqual(stale.status_code, 409)
        cancelled = self.mutate(
            self.owner,
            "DELETE",
            self.base + "/planned-payments/" + edited.json()["id"],
            None,
            version=edited.json()["version"],
        )
        self.assertEqual(cancelled.status_code, 200, cancelled.text)
        self.assertEqual(cancelled.json()["state"], "cancelled")
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                str(
                    conn.execute(
                        "SELECT total_amount FROM app.planned_payment_versions WHERE id=%s",
                        (p["id"],),
                    ).fetchone()[0]
                ),
                "123.45678",
            )


if __name__ == "__main__":
    suite = unittest.TestSuite(
        SettlementTests(name)
        for name in SettlementTests.__dict__
        if name.startswith("test_")
    )
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    (ROOT / "docs/settlements-validation.json").write_text(
        json.dumps(
            {
                "status": "passed"
                if result.wasSuccessful() and not result.skipped
                else "failed_or_not_run",
                "tests_run": result.testsRun,
                "failures": len(result.failures),
                "errors": len(result.errors),
                "skipped": len(result.skipped),
                "test_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                "api_source_sha256": {
                    str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in sorted((ROOT / "apps/api/didban").rglob("*.py"))
                },
            },
            indent=2,
        )
        + "\n"
    )
    raise SystemExit(0 if result.wasSuccessful() and not result.skipped else 1)
