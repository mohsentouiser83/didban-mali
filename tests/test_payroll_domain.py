"""Synthetic arithmetic fixtures; deliberately not a statutory-rate acceptance pack."""

import sys
from pathlib import Path
import unittest
import base64
from io import BytesIO
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps/api"))
from didban.domain.payroll.calculation import (
    calculate,
    PayrollValidation,
    validate_rules,
)
from didban.domain.payroll.validation import employee_profile, primary_sources
from didban.adapters.payroll_import import parse_import


def rules():
    return {
        "rounding_unit": "1",
        "monthly_reference_days": "30",
        "hours_per_day": "8",
        "overtime_multiplier": "1.4",
        "night_premium": "0.35",
        "holiday_multiplier": "1",
        "shift_rates": {
            "none": "0",
            "morning_evening": "0.1",
            "three_shift": "0.15",
            "night_rotation": "0.225",
        },
        "employee_insurance_rate": "0.07",
        "employer_insurance_rate": "0.23",
        "insurance_daily_floor": "0",
        "insurance_daily_cap": "10000",
        "tax_exemption_monthly": "500",
        "tax_bands": [{"up_to": None, "rate": "0.1"}],
        "tax_mode": "monthly",
        "eidi_wage_days": "60",
        "eidi_cap_daily_wage": "100",
        "sanavat_wage_days": "30",
        "year_days": "365",
        "annual_leave_days": "30",
        "special_treatment": {
            k: {"taxable": False, "insurable": False}
            for k in ("monthly", "eidi", "sanavat", "termination", "arrears")
        },
        "source_urls": [
            "https://natlex.ilo.org/dyn/natlex2/natlex2/files/download/21843/IRN21843%20Per.pdf"
        ],
    }


def contract():
    return {
        "employment_type": "full_time",
        "pay_basis": "monthly",
        "base_amount": "3000",
        "components": [],
    }


def work():
    return {
        "worked_days": "10",
        "worked_hours": "80",
        "period_days": "30",
        "shift_code": "none",
        "leave_opening_days": "2",
        "leave_accrued_days": "1",
        "leave_used_days": "1",
    }


class PayrollDomainTests(unittest.TestCase):
    def test_monthly_exact_components_and_cost(self):
        result = calculate(contract(), work(), rules())
        self.assertEqual(
            {
                k: result[k]
                for k in (
                    "gross",
                    "employee_insurance",
                    "employer_insurance",
                    "tax",
                    "net",
                    "employer_cost",
                    "leave_balance_days",
                )
            },
            {
                "gross": "1000",
                "employee_insurance": "70",
                "employer_insurance": "230",
                "tax": "50",
                "net": "880",
                "employer_cost": "1230",
                "leave_balance_days": "2",
            },
        )

    def test_hourly_and_part_time_not_double_prorated(self):
        c = contract() | {
            "pay_basis": "hourly",
            "base_amount": "10",
            "employment_type": "part_time",
            "hours_per_day": "4",
        }
        self.assertEqual(calculate(c, work(), rules())["gross"], "800")

    def test_missing_rule_or_tax_history_fails_closed(self):
        with self.assertRaises(PayrollValidation):
            validate_rules({})
        r = rules() | {"tax_mode": "cumulative"}
        with self.assertRaisesRegex(PayrollValidation, "missing_tax_history"):
            calculate(contract(), work(), r)

    def test_shift_night_overlap_rejected(self):
        with self.assertRaisesRegex(PayrollValidation, "night_and_shift_overlap"):
            calculate(
                contract(),
                work() | {"shift_code": "three_shift", "night_hours": "3"},
                rules(),
            )

    def test_eidi_prior_paid_and_sanavat(self):
        w = work() | {"service_days": "365", "prior_paid_special": "1000"}
        self.assertEqual(calculate(contract(), w, rules(), kind="eidi")["net"], "5000")
        self.assertEqual(
            calculate(contract(), w, rules(), kind="sanavat")["net"], "2000"
        )
        with self.assertRaisesRegex(
            PayrollValidation, "prior_paid_exceeds_entitlement"
        ):
            calculate(
                contract(), w | {"prior_paid_special": "7000"}, rules(), kind="eidi"
            )

    def test_termination_and_arrears_references(self):
        w = work() | {
            "service_days": "365",
            "leave_cashout_days": "2",
            "unpaid_wages": "100",
            "unpaid_wages_already_processed": True,
            "unpaid_wages_reference": "reference",
        }
        self.assertEqual(
            calculate(contract(), w, rules(), kind="termination")["net"], "3300"
        )
        with self.assertRaisesRegex(PayrollValidation, "missing_correction_reference"):
            calculate(
                contract(), work() | {"arrears_amount": "100"}, rules(), kind="arrears"
            )

    def test_loan_deductions_and_negative_leave_rejected(self):
        with self.assertRaisesRegex(PayrollValidation, "deductions_exceed_pay"):
            calculate(contract(), work() | {"loan_installment": "1000"}, rules())
        with self.assertRaisesRegex(PayrollValidation, "negative_leave_balance"):
            calculate(contract(), work() | {"leave_used_days": "9"}, rules())

    def test_opening_balances_need_date_and_source(self):
        with self.assertRaisesRegex(
            PayrollValidation, "opening_balance_source_required"
        ):
            employee_profile(
                {"name": "نمونه", "hired_on": "2026-01-01", "opening_loan": "100"}
            )
        self.assertFalse(primary_sources(["https://example.com/rules"]))
        self.assertTrue(primary_sources(rules()["source_urls"]))

    def test_csv_import_preserves_identifiers_and_rejects_duplicate(self):
        data = "personnel_code,name,national_id,hired_on\n0001,نمونه,0012345678,2026-01-01\n0001,تکراری,0012345678,2026-01-01\n".encode()
        _, preview = parse_import(
            {
                "filename": "employees.csv",
                "kind": "employees",
                "data_base64": base64.b64encode(data).decode(),
            }
        )
        self.assertEqual(preview["rows"][0]["personnel_code"], "0001")
        self.assertEqual(preview["rows"][0]["profile"]["national_id"], "0012345678")
        self.assertEqual(preview["errors"][0]["row"], 3)

    def test_summary_preserves_large_exact_aggregates(self):
        from didban.infrastructure.payroll import summary
        from decimal import Decimal

        value = "100000000000000000000000000000.123456"

        class Rows:
            def __init__(self, one=None, many=None):
                self.one, self.many = one, many

            def fetchone(self):
                return self.one

            def fetchall(self):
                return self.many

        class Conn:
            def execute(self, sql, params):
                if " AS day" in sql:
                    from datetime import date

                    return Rows({"day": date(2026, 9, 30)})
                if "AS hires" in sql:
                    return Rows({"hires": 2, "departures": 0})
                if "count(*) AS n" in sql:
                    return Rows({"n": 2})
                if "payroll_payment_links" in sql:
                    return Rows({"n": 0, "amount": Decimal(0)})
                return Rows(
                    many=[
                        {
                            "content": {
                                "totals": dict.fromkeys(
                                    ("gross", "net", "employer_cost"), value
                                )
                            }
                        }
                    ]
                    * 2
                )

        result = summary(Conn(), "fixture")
        self.assertEqual(result["net"], "200000000000000000000000000000.246912")

    def test_excel_payslip_headers_values_and_preview_metadata(self):
        from didban.adapters.payroll_export import export_run
        from openpyxl import load_workbook

        result = calculate(contract(), work(), rules())
        run = {
            "id": "synthetic",
            "company_name": "شرکت آزمون",
            "period_start": "2026-09-01",
            "period_end": "2026-09-30",
            "content": {
                "items": [
                    {
                        "personnel_code": "0001",
                        "profile": {"name": " =SUM(A1:A3)"},
                        "result": result,
                    }
                ]
            },
        }
        data, _, _ = export_run(run, "xlsx", None)
        book = load_workbook(BytesIO(data), data_only=False)
        rows = list(book["لیست حقوق"].values)
        self.assertEqual(len(rows[0]), len(rows[1]))
        values = dict(zip(rows[0], rows[1]))
        self.assertEqual(values["شرکت"], "شرکت آزمون")
        self.assertEqual(values["نام"], " =SUM(A1:A3)")
        self.assertEqual(values["خالص قابل پرداخت"], 880)
        self.assertEqual(book["لیست حقوق"]["E2"].number_format, "#,##0")
        self.assertEqual(book["لیست حقوق"]["C2"].data_type, "s")
        metadata = dict(book["مشخصات"].values)
        self.assertEqual(metadata["وضعیت"], "پیش‌نمایش؛ تأیید نشده")

    def test_import_money_units_are_explicit_and_toman_conversion_exact(self):
        def preview(unit):
            data = (
                "personnel_code,name,hired_on,opening_loan,opening_as_of,opening_source,unit\n0001,نمونه,2026-01-01,9007199254740993.123456,2026-08-31,مرجع,"
                + unit
                + "\n"
            ).encode()
            return parse_import(
                {
                    "kind": "employees",
                    "filename": "staff.csv",
                    "data_base64": base64.b64encode(data).decode(),
                }
            )[1]

        valid = preview("IRT")
        self.assertEqual(valid["errors"], [])
        self.assertEqual(
            valid["rows"][0]["profile"]["opening_loan"], "90071992547409931.23456"
        )
        self.assertEqual(valid["rows"][0]["profile"]["opening_unit_original"], "IRT")
        self.assertTrue(preview("")["errors"])
        self.assertTrue(preview("USD")["errors"])

    def test_import_preview_includes_failed_rows_error_columns_and_boolean_work(self):
        def preview(data, kind="employees"):
            return parse_import(
                {
                    "kind": kind,
                    "filename": "preview.csv",
                    "data_base64": base64.b64encode(data.encode()).decode(),
                }
            )[1]

        result = preview(
            "personnel_code,name,hired_on,opening_loan,unit\n001,,2026-01-01,0,IRR\n002,نمونه,2026-01-01,bad,IRT\n"
        )
        self.assertEqual(result["row_count"], 2)
        self.assertEqual(len(result["preview_rows"]), 2)
        self.assertEqual(
            [r["column"] for r in result["errors"]], ["name", "opening_loan"]
        )
        self.assertEqual(result["preview_rows"][1]["original"]["opening_loan"], "bad")
        self.assertIsNone(result["preview_rows"][0]["normalized"])
        valid = preview(
            "employee_id,period_start,period_end,worked_days,worked_hours,period_days,tax_history_complete,unpaid_wages_already_processed\n00000000-0000-0000-0000-000000000001,2026-09-01,2026-09-30,10,80,30,true,false\n",
            "work",
        )
        self.assertEqual(valid["errors"], [])
        self.assertTrue(valid["rows"][0]["content"]["tax_history_complete"])
        self.assertFalse(valid["rows"][0]["content"]["unpaid_wages_already_processed"])

    def test_work_xlsx_personnel_code_keeps_zero_prefix_and_rejects_numeric_cell(self):
        def parse(code):
            book = Workbook()
            sheet = book.active
            sheet.append(
                [
                    "personnel_code",
                    "employee_name",
                    "period_start",
                    "period_end",
                    "period_days",
                    "worked_days",
                    "worked_hours",
                    "deductions",
                    "unit",
                ]
            )
            sheet.append(
                [
                    code,
                    "نام نمونه",
                    "2026-09-01",
                    "2026-09-30",
                    "30",
                    "10",
                    "80",
                    "123.456789",
                    "IRT",
                ]
            )
            out = BytesIO()
            book.save(out)
            seen = []

            def resolve(employee, code):
                seen.append((employee, code))
                return {
                    "employee_id": "00000000-0000-0000-0000-000000000001",
                    "personnel_code": code,
                    "employee_name": "نام ثبت‌شده",
                }

            result = parse_import(
                {
                    "filename": "work.xlsx",
                    "kind": "work",
                    "data_base64": base64.b64encode(out.getvalue()).decode(),
                },
                resolve,
            )[1]
            return result, seen

        valid, seen = parse("0001")
        self.assertEqual(valid["errors"], [])
        self.assertEqual(seen, [("", "0001")])
        self.assertEqual(valid["preview_rows"][0]["original"]["personnel_code"], "0001")
        self.assertEqual(valid["rows"][0]["content"]["deductions"], "1234.56789")
        self.assertEqual(valid["rows"][0]["content"]["source_unit"], "IRT")
        invalid, seen = parse(1)
        self.assertEqual(seen, [])
        self.assertEqual(invalid["errors"][0]["code"], "identifier_must_be_text")
        self.assertEqual(invalid["errors"][0]["column"], "personnel_code")

    def test_special_tax_profile_is_exact_kind_and_preserves_regular_salary(self):
        r = rules()
        r["special_treatment"]["eidi"]["taxable"] = True
        r["special_tax_profiles"] = {
            "eidi": {
                "tax_mode": "monthly",
                "tax_exemption_monthly": "1000",
                "tax_bands": [{"up_to": None, "rate": "0.2"}],
            }
        }
        w = work() | {"service_days": "365"}
        special = calculate(contract(), w, r, kind="eidi")
        self.assertEqual(special["gross"], "6000")
        self.assertEqual(special["tax"], "1000")
        self.assertEqual(
            calculate(contract(), work(), r), calculate(contract(), work(), rules())
        )
        self.assertEqual(calculate(contract(), w, r, kind="sanavat")["tax"], "0")

    def test_special_cumulative_history_never_uses_regular_tax_history(self):
        r = rules()
        r["special_treatment"]["eidi"]["taxable"] = True
        r["special_tax_profiles"] = {
            "eidi": {
                "tax_mode": "cumulative",
                "tax_exemption_monthly": "1000",
                "tax_bands": [{"up_to": None, "rate": "0.2"}],
            }
        }
        w = work() | {
            "service_days": "365",
            "tax_history_complete": True,
            "tax_elapsed_months": "12",
            "prior_taxable": "999999",
            "prior_tax_paid": "999999",
        }
        with self.assertRaisesRegex(PayrollValidation, "missing_tax_history"):
            calculate(contract(), w, r, kind="eidi")
        w.update(
            {
                "eidi_tax_history_complete": True,
                "eidi_tax_elapsed_months": "1",
                "eidi_prior_taxable": "2000",
                "eidi_prior_tax_paid": "400",
            }
        )
        self.assertEqual(calculate(contract(), w, r, kind="eidi")["tax"], "1000")

    def test_import_special_tax_history_has_its_own_flags_and_exact_currency(self):
        data = (
            "employee_id,period_start,period_end,period_days,worked_days,worked_hours,eidi_tax_history_complete,eidi_tax_elapsed_months,eidi_prior_taxable,eidi_prior_tax_paid,unit\n"
            "00000000-0000-0000-0000-000000000001,2026-09-01,2026-09-30,30,10,80,true,1,123.456789,12.345678,IRT\n"
        ).encode()
        result = parse_import(
            {
                "kind": "work",
                "filename": "history.csv",
                "data_base64": base64.b64encode(data).decode(),
            }
        )[1]
        self.assertEqual(result["errors"], [])
        content = result["rows"][0]["content"]
        self.assertTrue(content["eidi_tax_history_complete"])
        self.assertEqual(content["eidi_prior_taxable"], "1234.56789")
        self.assertEqual(content["eidi_prior_tax_paid"], "123.45678")

    def test_special_tax_profile_validation_is_bounded_and_explicit(self):
        invalid = [
            None,
            {"monthly": {}},
            {"eidi": {}},
            {
                "eidi": {
                    "tax_mode": "monthly",
                    "tax_exemption_monthly": "0",
                    "tax_bands": [{"up_to": "100", "rate": "0.1"}],
                }
            },
            {
                "eidi": {
                    "tax_mode": "monthly",
                    "tax_exemption_monthly": "0",
                    "tax_bands": [{"up_to": None, "rate": "1.1"}],
                }
            },
        ]
        for profile in invalid:
            with self.subTest(profile=profile), self.assertRaises(PayrollValidation):
                validate_rules(rules() | {"special_tax_profiles": profile})

    def test_excel_formulas_rejected(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["personnel_code", "name", "hired_on"])
        sheet.append(["0001", "=1+2", "2026-01-01"])
        out = BytesIO()
        book.save(out)
        with self.assertRaisesRegex(PayrollValidation, "formulas_not_supported"):
            parse_import(
                {
                    "filename": "employees.xlsx",
                    "kind": "employees",
                    "data_base64": base64.b64encode(out.getvalue()).decode(),
                }
            )


if __name__ == "__main__":
    unittest.main()
