"""Offline exact/safe exports and real embedded Persian PDF rendering."""

import csv
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import Path
import sys
import unittest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps/api"))
from didban.adapters.report_rendering import render  # noqa: E402
from didban.domain.reports import digest, is_limited  # noqa: E402


def fixture():
    return {
        "schema_version": 1,
        "label": "تولیدشده سامانه",
        "manifest": {
            "id": "00000000-0000-4000-8000-000000000001",
            "company_name": "شرکت مرجع آزمایشی",
            "period_name": "شهریور ۱۴۰۵",
            "as_of_date": "2026-09-30",
            "manifest_hash": "a" * 64,
        },
        "metrics": [
            {
                "metric_key": "revenue",
                "basis": "ledger",
                "value": "123456789012345678901.123456",
                "status": "available",
                "reason_codes": [],
            },
            {
                "metric_key": "cash",
                "basis": "bank",
                "value": None,
                "status": "unavailable",
                "reason_codes": ["bank_anchor_unknown"],
            },
            {
                "metric_key": "expense",
                "basis": "ledger",
                "value": "0",
                "status": "available",
                "reason_codes": [],
            },
        ],
        "reconciliation_runs": [
            {"id": "run1", "kind": "bank_ledger", "missing_source_kinds": ["ledger"]}
        ],
        "candidates": [],
        "decisions": [],
        "cashflow": None,
        "settlement": None,
        "forecast": None,
        "sources": [],
        "rows": [
            {
                "source_file_id": "00000000-0000-4000-8000-000000000002",
                "sheet_name": '=HYPERLINK("http://evil.invalid")',
                "row_number": 4,
                "source_kind": "bank",
                "state": "rejected",
                "error_codes": ["invalid_amount"],
                "raw_values": {
                    "note": "=1+1",
                    "amount": "-12345678901234567890.123456",
                },
                "canonical_values": {"value": "-12345678901234567890.123456"},
            }
        ],
        "evidence": [],
    }


class RenderingTests(unittest.TestCase):
    def test_excel_exact_decimal_null_zero_and_no_formula(self):
        content = fixture()
        content["evidence"] = [
            {
                "id": "e1",
                "source_row_id": "r1",
                "rule_key": "long",
                "calculation": {"note": "x" * 65000},
            }
        ]
        data = render("xlsx", content, ROOT / "assets/fonts")
        book = load_workbook(BytesIO(data), data_only=False)
        values = [c.value for row in book["شاخص‌ها"] for c in row]
        self.assertIn("123,456,789,012,345,678,901.123456", values)
        self.assertIn(0, values)
        self.assertIn("unavailable", values)
        cells = [c for ws in book for row in ws for c in row]
        self.assertFalse(any(c.data_type == "f" for c in cells))
        self.assertTrue(
            any(c.value == '=HYPERLINK("http://evil.invalid")' for c in cells)
        )
        self.assertTrue(book["شاخص‌ها"].sheet_view.rightToLeft)
        chunks = [row[-1].value for row in list(book["شواهد"])[1:]]
        self.assertEqual(json.loads("".join(chunks))["note"], "x" * 65000)

    def test_toman_display_keeps_exact_fraction_and_snapshot_unit(self):
        content = fixture()
        content["manifest"]["display_unit"] = "IRT"
        data = render("xlsx", content, ROOT / "assets/fonts")
        book = load_workbook(BytesIO(data), data_only=False)
        values = [c.value for row in book["شاخص‌ها"] for c in row]
        self.assertIn("12,345,678,901,234,567,890.1123456", values)
        self.assertEqual(content["metrics"][0]["value"], "123456789012345678901.123456")

    def test_excel_grouped_numeric_cells_and_large_exact_text(self):
        content = fixture()
        content["metrics"][2]["value"] = "1234567.125000"
        book = load_workbook(BytesIO(render("xlsx", content, ROOT / "assets/fonts")))
        # A is the continuation marker; D is the financial amount.
        small = book["شاخص‌ها"]["D4"]
        self.assertEqual(small.value, 1234567.125)
        self.assertEqual(small.data_type, "n")
        self.assertEqual(small.number_format, "#,##0.000000")
        large = book["شاخص‌ها"]["D2"]
        self.assertEqual(large.data_type, "s")
        self.assertEqual(large.value.replace(",", ""), content["metrics"][0]["value"])

    def test_error_csv_is_sanitized_and_includes_original_row(self):
        raw = render("error_csv", fixture(), ROOT / "assets/fonts").decode("utf-8-sig")
        rows = list(csv.reader(StringIO(raw)))
        self.assertEqual(len(rows), 3)
        self.assertTrue(rows[2][4].startswith("'="))
        self.assertEqual(
            json.loads(rows[2][8])["amount"], "-12345678901234567890.123456"
        )

    def test_real_errors_and_overlap_exclusions_have_separate_excel_sheets_and_csv_types(
        self,
    ):
        from copy import deepcopy

        content = fixture()
        excluded = deepcopy(content["rows"][0])
        excluded.update(row_number=5, state="excluded_overlap", error_codes=[])
        valid = deepcopy(excluded)
        valid.update(row_number=6, state="valid")
        conflicting = deepcopy(excluded)
        conflicting.update(row_number=7, error_codes=["invalid_amount"])
        content["rows"] += [excluded, valid, conflicting]
        original = deepcopy(content)
        rows = list(
            csv.DictReader(
                StringIO(
                    render("error_csv", content, ROOT / "assets/fonts").decode(
                        "utf-8-sig"
                    )
                )
            )
        )
        self.assertEqual(
            [row["record_type"] for row in rows],
            ["report", "error", "excluded_overlap", "error"],
        )
        self.assertEqual(rows[2]["errors"], "")
        self.assertEqual(rows[2]["classification"], "تکراریِ کنارگذاشته‌شده")
        self.assertTrue(rows[2]["reason"])
        book = load_workbook(BytesIO(render("xlsx", content, ROOT / "assets/fonts")))
        self.assertNotIn("ردیف‌های خطادار", book.sheetnames)
        errors = list(book["خطاهای واقعی"].values)
        exclusions = list(book["ردیف‌های کنارگذاشته‌شده"].values)
        self.assertEqual([row[3] for row in errors[1:]], ["4", "7"])
        self.assertEqual([row[3] for row in exclusions[1:]], ["5"])
        self.assertEqual(content, original)

    def test_duplicate_only_bundle_has_no_real_errors(self):
        content = fixture()
        content["rows"][0].update(state="excluded_overlap", error_codes=[])
        book = load_workbook(BytesIO(render("xlsx", content, ROOT / "assets/fonts")))
        self.assertEqual(book["خطاهای واقعی"].max_row, 1)
        self.assertEqual(book["ردیف‌های کنارگذاشته‌شده"].max_row, 2)

    def test_pdf_embeds_font_and_can_be_rendered(self):
        data = render("pdf", fixture(), ROOT / "assets/fonts")
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertIn(b"/FontFile2", data)
        out = ROOT / ".local/report-validation"
        out.mkdir(parents=True, exist_ok=True)
        (out / "managerial.pdf").write_bytes(data)

    def test_report_digest_changes_for_new_pinned_decision(self):
        content = fixture()
        before = digest(content)
        content["decisions"].append({"id": "new-decision"})
        self.assertNotEqual(before, digest(content))
        self.assertTrue(is_limited(content))


if __name__ == "__main__":
    result = unittest.TextTestRunner(verbosity=2).run(
        unittest.defaultTestLoader.loadTestsFromTestCase(RenderingTests)
    )
    (ROOT / "docs/report-rendering-validation.json").write_text(
        json.dumps(
            {
                "status": "passed" if result.wasSuccessful() else "failed",
                "tests_run": result.testsRun,
                "failures": len(result.failures),
                "errors": len(result.errors),
                "live_model_calls": 0,
                "test_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            },
            indent=2,
        )
        + "\n"
    )
    raise SystemExit(0 if result.wasSuccessful() else 1)
