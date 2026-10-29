"""Isolated HTTP/worker acceptance for portal files and versioned automatic retry."""

import io
import hashlib
import json
import os
from pathlib import Path
from decimal import Decimal
import unittest
from unittest.mock import patch
from uuid import uuid4

import openpyxl
import psycopg

from test_interpretation import InterpretationTests
from test_moadian_export import portal_sheet


def portal_workbook():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "invoice"
    for row in portal_sheet().rows:
        sheet.append(row.values)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class PortalRuntimeTests(InterpretationTests):
    def verify_supplied_exports(self):
        """Opt-in only: unchanged real files in a disposable company/database."""
        directory = Path(os.environ["DIDBAN_MOADIAN_REAL_SOURCE_DIR"])
        expected = {7: 25, 8: 44, 9: 107, 10: 104, 11: 7, 12: 37, 13: 67, 14: 33}
        session = self.session(
            self.period(starts_on="2024-03-20", ends_on="2025-03-20")
        )
        hashes = {}
        expected_amounts = {}
        for number in expected:
            path = directory / f"report ({number}).xlsx"
            content = path.read_bytes()
            hashes[path.name] = hashlib.sha256(content).hexdigest()
            book = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=False
            )
            try:
                source_rows = list(book.active.iter_rows(min_row=2, values_only=True))
                expected_amounts[path.name] = (
                    sum(Decimal(str(row[5])) for row in source_rows),
                    sum(Decimal(str(row[6])) for row in source_rows),
                )
            finally:
                book.close()
            result = self.upload(session, content, name=path.name)
            self.assertEqual(result.status_code, 202, result.text)
        self.assertEqual(self.seal(session).status_code, 202)
        self.assertIsNotNone(
            self.complete_coverage(session)["interpretation_finished_at"]
        )
        facts = []
        for file in self.files(session):
            number = int(file["original_name"].split("(")[1].split(")")[0])
            summary = file["interpretation_summary"]
            self.assertEqual(summary["valid_rows"], expected[number])
            self.assertEqual(summary["rejected_rows"], 0)
            self.assertEqual(summary["excluded_rows"], 0)
            source = summary["sheets"][0]
            self.assertEqual(
                source["source_kind"],
                "moadian_purchases" if number <= 10 else "moadian_sales",
            )
            self.assertEqual(source["provenance"], "validated_automatic")
            normalized = self.owner.get(
                self.file_path(session, file) + "/rows?limit=100"
            ).json()
            normalized_rows = normalized["items"]
            while normalized.get("next_cursor"):
                normalized = self.owner.get(
                    self.file_path(session, file) + "/rows",
                    params={"limit": 100, "cursor": normalized["next_cursor"]},
                ).json()
                normalized_rows.extend(normalized["items"])
            amounts = [
                {value["field"]: value["value"] for value in row["values"]}
                for row in normalized_rows
            ]
            self.assertEqual(
                (
                    sum(Decimal(value["gross"]) for value in amounts),
                    sum(Decimal(value["vat"]) for value in amounts),
                ),
                expected_amounts[file["original_name"]],
            )
            facts.append(
                {
                    "file": file["original_name"],
                    "rows": summary["valid_rows"],
                    "kind": source["source_kind"],
                }
            )
        for filename, digest in hashes.items():
            self.assertEqual(
                hashlib.sha256((directory / filename).read_bytes()).hexdigest(), digest
            )
        with psycopg.connect(self.dsn) as conn:
            calls = conn.execute(
                "SELECT count(*) FROM app.model_requests WHERE company_id=%s",
                (self.company,),
            ).fetchone()[0]
            self.assertEqual(calls, 0)
        output = Path(__file__).resolve().parents[1] / ".local/moadian-validation"
        output.mkdir(parents=True, exist_ok=True)
        (output / "real-http.json").write_text(
            json.dumps(
                {
                    "status": "passed",
                    "scope": "isolated_HTTP_private_real_files",
                    "files": facts,
                    "valid_rows": sum(item["rows"] for item in facts),
                    "model_calls": calls,
                    "original_files_unchanged": True,
                    "canonical_amounts_match_original": True,
                    "coverage_answer": "unknown_in_disposable_test_only",
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n"
        )

    def prepared(self):
        session = self.session(
            self.period(starts_on="2024-03-20", ends_on="2025-03-20")
        )
        self.assertEqual(
            self.upload(session, portal_workbook(), name="export.xlsx").status_code, 202
        )
        self.assertEqual(self.seal(session).status_code, 202)
        return session

    def test_portal_upload_uses_tax_identity_and_does_not_create_book_invoices(self):
        session = self.prepared()
        state = self.complete_coverage(session)
        self.assertIsNotNone(state["interpretation_finished_at"])
        file = self.files(session)[0]
        self.assertEqual(file["interpretation_summary"]["valid_rows"], 1)
        self.assertEqual(file["interpretation_summary"]["rejected_rows"], 0)
        summary = file["interpretation_summary"]["sheets"][0]
        self.assertEqual(summary["source_kind"], "moadian_sales")
        self.assertEqual(summary["provenance"], "validated_automatic")
        rows = self.owner.get(self.file_path(session, file) + "/rows").json()["items"]
        values = {v["field"]: v["value"] for v in rows[0]["values"]}
        self.assertIsNone(values["invoice_id"])
        self.assertEqual(values["tax_id"], "TAX0001")
        self.assertEqual(rows[0]["raw_values"][6], "1403/12/28 23:59:59")
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT count(*) FROM app.model_requests WHERE company_id=%s",
                    (self.company,),
                ).fetchone()[0],
                0,
            )

    def test_retry_supersedes_mapping_only_and_is_versioned_idempotent_and_authorized(
        self,
    ):
        session = self.prepared()
        with patch("didban.domain.interpretation.portal_export", return_value=None):
            self.process(session)
        issue = self.issue(session)
        self.assertEqual(issue["code"], "interpretation_mapping")
        version = self.detail(session)["session_version"]
        path = self.base + "/processing-sessions/" + session["id"] + "/interpret"
        self.assertEqual(
            self.mutate(
                self.client("viewer"), "POST", path, version=version
            ).status_code,
            403,
        )
        self.assertEqual(
            self.mutate(self.owner, "POST", path, version=version - 1).status_code, 409
        )
        key = uuid4().hex
        first = self.mutate(self.owner, "POST", path, version=version, key=key)
        self.assertEqual(first.status_code, 202, first.text)
        self.assertEqual(
            self.mutate(self.owner, "POST", path, version=version, key=key).json(),
            first.json(),
        )
        # Reprocessing is automatic; it must not manufacture a coverage answer.
        self.process(session)
        self.assertEqual(self.issue(session)["code"], "interpretation_coverage")
        with psycopg.connect(self.dsn) as conn:
            old = conn.execute(
                "SELECT superseded_at,resolved_at FROM app.processing_issues WHERE company_id=%s AND id=%s",
                (self.company, issue["id"]),
            ).fetchone()
            self.assertIsNotNone(old[0])
            self.assertIsNone(old[1])
            self.assertEqual(
                conn.execute(
                    "SELECT count(*) FROM app.issue_answers WHERE company_id=%s",
                    (self.company,),
                ).fetchone()[0],
                0,
            )
        self.assertIsNotNone(
            self.complete_coverage(session)["interpretation_finished_at"]
        )
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT count(*) FROM app.interpretation_runs WHERE company_id=%s",
                    (self.company,),
                ).fetchone()[0],
                1,
            )

    def test_coverage_question_cannot_be_cleared_by_automatic_retry(self):
        session = self.prepared()
        self.process(session)
        self.assertEqual(self.issue(session)["code"], "interpretation_coverage")
        path = self.base + "/processing-sessions/" + session["id"] + "/interpret"
        result = self.mutate(
            self.owner, "POST", path, version=self.detail(session)["session_version"]
        )
        self.assertEqual(result.status_code, 409)
        self.assertEqual(self.detail(session)["status"], "waiting_input")

    def test_retry_preserves_other_required_barriers_and_worker_cannot_pass_them(self):
        session = self.prepared()
        with patch("didban.domain.interpretation.portal_export", return_value=None):
            self.process(session)
        issue = self.issue(session)
        with psycopg.connect(self.dsn) as conn:
            conn.execute("SET LOCAL search_path=app,pg_catalog")
            other = conn.execute(
                """INSERT INTO app.processing_issues(
                company_id,session_id,code,root_key,required,question_fa,answer_kind,answer_constraints)
                SELECT company_id,session_id,'interpretation_unit',root_key||':unrelated',
                true,'واحد منبع دیگر چیست؟',answer_kind,answer_constraints
                FROM app.processing_issues WHERE company_id=%s AND id=%s RETURNING id""",
                (self.company, issue["id"]),
            ).fetchone()[0]
        path = self.base + "/processing-sessions/" + session["id"] + "/interpret"
        result = self.mutate(
            self.owner, "POST", path, version=self.detail(session)["session_version"]
        )
        self.assertEqual(result.status_code, 202, result.text)
        self.assertEqual(self.detail(session)["status"], "waiting_input")
        self.assertFalse(self.interpreter().run_once())
        with psycopg.connect(self.dsn) as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT resolved_at,superseded_at FROM app.processing_issues WHERE company_id=%s AND id=%s",
                    (self.company, other),
                ).fetchone(),
                (None, None),
            )


if __name__ == "__main__":
    suite = unittest.TestSuite(
        PortalRuntimeTests(name)
        for name in sorted(PortalRuntimeTests.__dict__)
        if name.startswith("test_")
    )
    if os.environ.get("DIDBAN_MOADIAN_REAL_SOURCE_DIR"):
        suite.addTest(PortalRuntimeTests("verify_supplied_exports"))
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(0 if result.wasSuccessful() and not result.skipped else 1)
