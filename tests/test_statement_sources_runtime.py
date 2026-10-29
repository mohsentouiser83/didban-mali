"""Disposable PostgreSQL: independent TB ingestion, barriers and pinned exports."""

import unittest
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from test_reports_runtime import ReportTests
from test_statement_sources import trial_workbook, cash_rows
from didban.application.financial_worker import FinancialWorker


class StatementSourceRuntimeTests(ReportTests):
    def publish_source(self, data, name):
        period = self.period(starts_on="2026-09-01", ends_on="2026-09-30")
        session = self.session(period)
        result = self.upload(session, data, name=name)
        self.assertEqual(result.status_code, 202, result.text)
        self.assertEqual(self.seal(session).status_code, 202)
        self.complete_coverage(session)
        self.start_financial(session)
        worker = FinancialWorker(self.worker_settings)
        for _ in range(40):
            state = self.detail(session)
            if state.get("manifest_id"):
                break
            if state["status"] == "waiting_input":
                issue = self.issue(session)
                choices = issue["constraints"]["choices"]
                key = issue["constraints"].get("financial_key", "")
                value = (
                    "recognized"
                    if key.startswith("income_tax|") and name == "trial.xlsx"
                    else next(
                        c
                        for c in [
                            "declared_complete",
                            "affirmed",
                            "zero_confirmed",
                            "not_required",
                            "unknown",
                        ]
                        if c in choices
                    )
                )
                answered = self.answer(session, value)
                self.assertEqual(answered.status_code, 201, answered.text)
            self.assertTrue(worker.run_once())
        else:
            self.fail("Source did not publish")
        self.start_analysis(session)
        self.finish_analysis(session)
        reports = self.owner.get(self.base + "/reports").json()["items"]
        return session, reports[-1]

    def test_trial_source_and_exports_without_fabricated_journals(self):
        session, report = self.publish_source(trial_workbook(), "trial.xlsx")
        self.assertIsNone(self.metrics(report["manifest_id"])["revenue"]["value"])
        url = self.base + "/reports/" + report["id"] + "/management"
        result = self.owner.get(url)
        self.assertEqual(result.status_code, 200, result.text)
        data = result.json()
        self.assertEqual(data["statement_basis"], "trial_balance")
        self.assertEqual(
            next(t for t in data["statements"]["totals"] if t["key"] == "net_profit")[
                "value"
            ],
            "250",
        )
        self.assertEqual(
            next(t for t in report["metrics"] if t["metric_key"] == "net_profit")[
                "basis"
            ],
            "trial_balance",
        )
        self.assertEqual(
            next(t for t in report["metrics"] if t["metric_key"] == "net_profit")[
                "value"
            ],
            "250",
        )
        self.assertEqual(data["cashflow_statement"]["status"], "unavailable")
        self.assertEqual(self.client("outsider").get(url).status_code, 404)
        self.drain()
        exact = self.owner.get(self.base + "/reports/" + report["id"]).json()
        self.assertTrue(all(job["status"] == "ready" for job in exact["exports"]))
        proof = Path(".local/report-validation/statements-v2")
        proof.mkdir(parents=True, exist_ok=True)
        for job in exact["exports"]:
            content = self.owner.get(
                self.base + "/exports/" + job["id"] + "/content"
            ).content
            if job["format"] in ("pdf", "xlsx"):
                (proof / ("trial-report." + job["format"])).write_bytes(content)
            if job["format"] == "xlsx":
                book = load_workbook(BytesIO(content))
                self.assertIn("صورت جریان وجوه نقد", book.sheetnames)
        original = data
        revised = self.mutate(
            self.owner,
            "POST",
            self.base + "/results/" + report["manifest_id"] + "/management-reports",
            {"as_of_date": "2026-09-15"},
        )
        self.assertEqual(revised.status_code, 201, revised.text)
        shortened = self.owner.get(
            self.base + "/reports/" + revised.json()["id"] + "/management"
        ).json()
        self.assertTrue(
            all(line["value"] is None for line in shortened["statements"]["lines"])
        )
        self.assertEqual(self.owner.get(url).json(), original)
        template = self.owner.get(
            self.base + "/reports/" + report["id"] + "/statement-template"
        )
        book = load_workbook(BytesIO(template.content))
        self.assertTrue(
            any(cell.value == "sales" for row in book.active for cell in row)
        )

    def test_trial_metadata_questions_pause_and_resume_with_audited_settings(self):
        from test_statement_sources import FIELDS

        period = self.period(starts_on="2026-09-01", ends_on="2026-09-30")
        session = self.session(period)
        uploaded = self.upload(
            session, trial_workbook(False), name="trial-no-dates.xlsx"
        )
        self.assertEqual(uploaded.status_code, 202, uploaded.text)
        self.assertEqual(self.seal(session).status_code, 202)
        self.complete_coverage(session)
        issue = self.issue(session)
        self.assertEqual(issue["code"], "interpretation_mapping")
        self.assertEqual(issue["constraints"]["setting_key"], "mapping")
        self.assertFalse(self.interpreter().run_once())
        self.assertEqual(self.issue(session)["id"], issue["id"])
        file = self.files(session)[0]
        path = self.file_path(session, file) + "/interpretation"
        settings = {
            "sheet_name": "تراز",
            "header_row": 1,
            "source_kind": "trial_balance",
            "unit": None,
            "calendar": "gregorian",
            "period_start": "2026-09-01",
            "as_of_date": "2026-09-30",
            "mapping": [
                {"target_field": field, "source_column": str(index + 1)}
                for index, field in enumerate(FIELDS[:8] + ["currency"])
            ],
            "reason": "تعیین بازه تراز مرجع",
        }
        configured = self.mutate(
            self.owner,
            "PUT",
            path,
            settings,
            version=self.detail(session)["session_version"],
        )
        self.assertEqual(configured.status_code, 202, configured.text)
        self.process(session)
        self.assertEqual(
            self.issue(session)["constraints"]["setting_key"], "trial_balance_level"
        )
        answered = self.answer(session, "detail")
        self.assertEqual(answered.status_code, 201, answered.text)
        self.process(session)
        self.assertEqual(
            self.issue(session)["constraints"]["setting_key"], "trial_balance_flows"
        )
        answered = self.answer(session, "before_closing")
        self.assertEqual(answered.status_code, 201, answered.text)
        state = self.complete_coverage(session)
        self.assertIsNotNone(state["interpretation_finished_at"])
        preview = self.owner.get(path).json()["previews"][0]
        self.assertEqual(preview["period_start"], "2026-09-01")
        self.assertEqual(preview["as_of_date"], "2026-09-30")
        self.assertEqual(
            self.files(session)[0]["interpretation_summary"]["valid_rows"], 6
        )

    def test_classified_cash_upload_reaches_report(self):
        import csv
        from io import StringIO

        stream = StringIO()
        fields = [
            "entry_id",
            "entry_type",
            "date",
            "account",
            "role",
            "debit",
            "credit",
            "unit",
            "cashflow_activity",
        ]
        writer = csv.writer(stream)
        writer.writerow([("currency" if key == "unit" else key) for key in fields])
        for row in cash_rows():
            writer.writerow([row.values.get(key, "") for key in fields])
        session, report = self.publish_source(stream.getvalue().encode(), "ledger.csv")
        cash = self.owner.get(
            self.base + "/reports/" + report["id"] + "/management"
        ).json()["cashflow_statement"]
        self.assertEqual(cash["status"], "available")
        self.assertEqual(cash["closing"]["value"], "1400")
        self.assertEqual(
            [s["net"]["value"] for s in cash["sections"]], ["300", "-300", "400"]
        )
        self.drain()
        exact = self.owner.get(self.base + "/reports/" + report["id"]).json()
        self.assertTrue(all(j["status"] == "ready" for j in exact["exports"]))
        proof = Path(".local/report-validation/statements-v2")
        proof.mkdir(parents=True, exist_ok=True)
        for job in exact["exports"]:
            if job["format"] in ("pdf", "xlsx"):
                (proof / ("cash-report." + job["format"])).write_bytes(
                    self.owner.get(
                        self.base + "/exports/" + job["id"] + "/content"
                    ).content
                )


if __name__ == "__main__":
    suite = unittest.TestSuite(
        StatementSourceRuntimeTests(name)
        for name in [
            "test_trial_source_and_exports_without_fabricated_journals",
            "test_classified_cash_upload_reaches_report",
            "test_trial_metadata_questions_pause_and_resume_with_audited_settings",
        ]
    )
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(not result.wasSuccessful())
