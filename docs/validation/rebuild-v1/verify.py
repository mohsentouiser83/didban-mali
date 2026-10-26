"""Verify saved synthetic fixtures, never import or exercise the application engine.

Requires Python 3 and openpyxl (read-only use). Golden JSON values are literals
authored separately; this verifier does not replace a future application adapter.
"""
from __future__ import annotations

import csv
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, getcontext
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
D = Decimal
getcontext().prec = 38
CHECKS: list[dict] = []


def read_json(name: str):
    return json.loads((ROOT / name).read_text())


def write_json(name: str, value):
    (ROOT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def check(name: str, actual, expected):
    if actual != expected:
        raise AssertionError(f'{name}: actual={actual!r}, expected={expected!r}')
    CHECKS.append({'check': name, 'status': 'passed'})


def xrows(name: str, sheet: str):
    workbook = load_workbook(ROOT / name, read_only=True, data_only=True)
    rows = list(workbook[sheet].values)
    workbook.close()
    return [dict(zip(rows[0], [v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else v for v in row])) for row in rows[1:]]


def crows(name: str):
    with (ROOT / name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def sumd(values):
    return sum((D(str(v)) for v in values), D(0))


def baseline():
    golden = read_json('expected/baseline.json')
    ledger = xrows('inputs/baseline/ledger.xlsx', 'دفتر')
    banks = {key: xrows('inputs/baseline/bank.xlsx', sheet) for key, sheet in [('bank_a', 'بانک الف'), ('bank_b', 'بانک ب')]}
    invoices = {key: xrows('inputs/baseline/invoices.xlsx', sheet) for key, sheet in [('sales', 'فروش'), ('purchase', 'خرید')]}
    payments = crows('inputs/baseline/settlements.csv')
    by_entry = defaultdict(list)
    for row in ledger:
        by_entry[row['entry_id']].append(row)
        date.fromisoformat(row['date'])
        check('unit:' + row['row_id'], row['currency'], 'IRR')
    for entry, rows in by_entry.items():
        check('balanced:' + entry, sumd(r['debit'] for r in rows), sumd(r['credit'] for r in rows))
    check('ledger text ID', ledger[0]['row_id'], '00001')
    check('bank text ID', banks['bank_a'][0]['row_id'], '00001')
    date_book = load_workbook(ROOT / 'inputs/baseline/bank.xlsx', read_only=True)
    check('real Excel date', isinstance(date_book['بانک الف']['B2'].value, datetime), True)
    date_book.close()
    for cutoff, start, field in [('2026-08-31', '2026-08-01', 'previous'), ('2026-09-30', '2026-09-01', 'current')]:
        period = [r for r in ledger if start <= r['date'] <= cutoff and r['entry_type'] != 'opening']
        all_rows = [r for r in ledger if r['date'] <= cutoff]
        income = sumd(D(str(r['credit'])) - D(str(r['debit'])) for r in period if r['role'] == 'revenue')
        cost = sumd(D(str(r['debit'])) - D(str(r['credit'])) for r in period if r['role'] in {'cogs', 'operating_expense', 'bank_fee', 'payroll_expense', 'income_tax_expense'})
        values = {'revenue': income, 'expense': cost, 'net_profit': income - cost}
        values['ar_ledger'] = sumd(D(str(r['debit'])) - D(str(r['credit'])) for r in all_rows if r['role'] == 'trade_ar')
        values['ap_trade_ledger'] = sumd(D(str(r['credit'])) - D(str(r['debit'])) for r in all_rows if r['role'] == 'trade_ap')
        values['debt_total_ledger'] = sumd(D(str(r['credit'])) - D(str(r['debit'])) for r in all_rows if r['role'].endswith('_payable') or r['role'] == 'trade_ap')
        values['cash'] = D(0)
        for key, rows in banks.items():
            eligible = [r for r in rows if r['date'] <= cutoff]
            balance = D(str(eligible[0]['balance']))
            for row in eligible[1:]:
                balance += D(str(row['amount']))
                check(f'running balance:{key}:{row["row_id"]}:{field}', balance, D(str(row['balance'])))
            ledger_bank = sumd(D(str(r['debit'])) - D(str(r['credit'])) for r in all_rows if r['role'] == key)
            check(f'ledger-bank reconciliation:{key}:{field}', balance, ledger_bank)
            values['cash'] += balance
            if field == 'current':
                check('bank golden:' + key, balance, D(golden['bank_balances'][key]))
        for direction, metric in [('sales', 'ar_invoice'), ('purchase', 'ap_invoice')]:
            total = D(0)
            for inv in invoices[direction]:
                if inv['issued_on'] > cutoff:
                    continue
                paid = sumd(p['amount'] for p in payments if p['invoice_id'] == inv['invoice_id'] and p['date'] <= cutoff)
                remaining = D(str(inv['gross'])) - paid
                total += remaining
                if field == 'current':
                    check('invoice remaining:' + inv['invoice_id'], remaining, D(golden['invoice_remaining'][inv['invoice_id']]))
            values[metric] = total
        for key, value in values.items():
            check(f'golden metric:{key}:{field}', value, D(golden['metrics'][key][field]))
    for key, m in golden['metrics'].items():
        current, previous = D(m['current']), D(m['previous'])
        check('delta:' + key, current - previous, D(m['delta']))
        check('percent:' + key, round((current - previous) / previous * 100, 24), round(D(m['percent']), 24))
    check('AR comparable difference', D(golden['metrics']['ar_ledger']['current']) - D(golden['metrics']['ar_invoice']['current']), D(golden['comparable_difference']['ar']))
    check('AP comparable difference', D(golden['metrics']['ap_trade_ledger']['current']) - D(golden['metrics']['ap_invoice']['current']), D(golden['comparable_difference']['ap']))
    for direction in ('sales', 'purchase'):
        portal = crows('inputs/baseline/moadian-' + ('purchases' if direction == 'purchase' else 'sales') + '.csv')
        internal = {r['tax_id']: r for r in invoices[direction] if '2026-09-01' <= r['issued_on'] <= '2026-09-30'}
        matches, differences, only = [], [], []
        for p in portal:
            inv = internal.get(p['tax_id'])
            if inv is None:
                only.append(p['invoice_id'])
            elif D(p['gross']) == D(str(inv['gross'])) and D(p['vat']) == D(str(inv['vat'])):
                matches.append(inv['invoice_id'])
            else:
                differences.append({'invoice_id': inv['invoice_id'], 'internal': str(inv['gross']), 'portal': p['gross'], 'difference': str(D(p['gross']) - D(str(inv['gross'])))})
        check('moadian matches:' + direction, matches, golden['moadian'][direction]['matched'])
        if direction == 'sales':
            check('moadian discrepancies', differences, golden['moadian']['sales']['amount_mismatch'])
            check('moadian portal only', only, golden['moadian']['sales']['portal_only'])
    current_movements = [r for rows in banks.values() for r in rows if '2026-09-01' <= r['date'] <= '2026-09-30' and r['row_type'] == 'movement']
    for key, internal, positive in [('external_in', False, True), ('external_out', False, False), ('internal_in', True, True), ('internal_out', True, False)]:
        result = sumd(abs(D(str(r['amount']))) for r in current_movements if (r['reference'] == 'TRANSFER1') == internal and (D(str(r['amount'])) > 0) == positive)
        check('cash movement:' + key, result, D(golden['cash_movements'][key]))
    history = crows('inputs/baseline/receipt-history.csv')
    lag_weights = sorted(((date.fromisoformat(r['paid_on']) - date.fromisoformat(r['due_on'])).days, D(r['amount'])) for r in history if r['paid_on'] <= '2026-09-30')
    threshold = sumd(w for _, w in lag_weights) / 2
    cumulative = D(0)
    lag = None
    for days, weight in lag_weights:
        cumulative += weight
        if cumulative >= threshold:
            lag = days
            break
    check('weighted median delay', lag, golden['forecast']['weighted_median_delay_days'])
    check('overdue forecast requires answer', date.fromisoformat('2026-09-20') + timedelta(days=lag) < date.fromisoformat('2026-09-30'), True)
    context = read_json('scenarios/baseline-context.json')
    answers = {a['issue_id']: a['value'] for a in context['answers']}
    check('question barrier before answers', golden['initial_workflow']['publish_allowed'], False)
    check('two explicit answers', len(answers), 2)
    source_invoice = next(r for r in invoices['sales'] if r['invoice_id'] == 'S001')
    remaining_s001 = D(str(source_invoice['gross'])) - sumd(p['amount'] for p in payments if p['invoice_id'] == 'S001' and p['date'] <= '2026-09-30')
    forecast_events = [(answers['Q-RECEIPT-S001'], 'receipt', 'S001', remaining_s001)]
    plans = crows('inputs/baseline/planned-payments.csv')
    for p in plans:
        if p['confirmed'] == 'true' and '2026-10-01' <= p['date'] <= '2026-10-30':
            remaining = D(p['total_obligation']) - D(p['paid_as_of'])
            if p['invoice_id']:
                check('plan-invoice same obligation:' + p['plan_id'], remaining, D(golden['invoice_remaining'][p['invoice_id']]))
            forecast_events.append((p['date'], 'payment', p['obligation_id'], remaining))
    forecast_events.sort()
    balance, receipts, outflow = values['cash'], D(0), D(0)
    events = []
    daily = []
    for day in range(1, 31):
        on = (date(2026, 9, 30) + timedelta(days=day)).isoformat()
        for when, direction, obligation, amount in forecast_events:
            if when != on:
                continue
            if direction == 'receipt':
                receipts += amount
                balance += amount
            else:
                outflow += amount
                balance -= amount
            events.append({'date': when, 'direction': direction, 'obligation_id': obligation, 'amount': str(amount), 'balance': str(balance)})
        daily.append({'date': on, 'balance': str(balance), 'status': 'estimated'})
    check('forecast events', events, golden['forecast']['events'])
    for key, value in [('closing', balance), ('receipts', receipts), ('payments', outflow)]:
        check('forecast:' + key, value, D(golden['forecast'][key]))
    write_json('expected/forecast-daily.json', {'as_of_date': '2026-09-30', 'status': 'estimated', 'rows': daily})
    reference = load_workbook(ROOT / 'expected/reference-results.xlsx', read_only=True, data_only=True)
    for row in list(reference['خلاصه'].values)[1:]:
        key = row[5]
        for i, field in enumerate(('current', 'previous', 'delta'), start=1):
            check('workbook reference:' + key + ':' + field, D(str(row[i])), D(golden['metrics'][key][field]))
    reference.close()


def domain_cases():
    """Independent arithmetic/spec checks, not actual domain service tests."""
    for i in range(1, 16):
        s = read_json(f'scenarios/F{i:02}.json')
        x, e, key = s['input'], s['expected'], s['id']
        if i == 1:
            check(key, D(x['amount']) * 10, D(e['amount_irr']))
        elif i == 2:
            check(key + ':revenue', sumd(r['amount'] for r in x['ledger'] if r['role'] == 'revenue'), D(e['revenue']))
            check(key + ':bank', sumd(x['bank_in']), D(e['bank_in']))
        elif i == 3:
            check(key, D(x['revenue']) - D(x['expense_before_tax']) - D(x['tax']), D(e['net_profit']))
            check(key + ':unknown', e['missing_tax_variant']['net_profit'], None)
        elif i == 4:
            check(key + ':income', sumd(r.get('credit', '0') for r in x['lines'] if r['kind'] == 'regular'), D(e['revenue']))
            check(key + ':adjustment', sumd(r.get('debit', '0') for r in x['lines'] if r['kind'] == 'real_cogs_adjustment'), D(e['cogs']))
        elif i == 5:
            check(key, D(x['anchor']['balance']) + sumd(r['amount'] for r in x['movements'] if x['anchor']['date'] < r['date'] <= x['as_of']), D(e['balance']))
        elif i == 6:
            check(key, sumd(r['amount'] for r in x['movements']), D(e['total_net']))
            check(key + ':internal', sumd(r['amount'] for r in x['movements'] if r['type'] == 'internal'), D(e['internal_net']))
        elif i == 7:
            for cutoff, field in [(x['cutoff'], 'period_remaining'), (x['current'], 'current_remaining')]:
                check(key + field, D(x['gross']) - sumd(p['amount'] for p in x['payments'] if p['date'] <= cutoff), D(e[field]))
        elif i == 8:
            paid = D(x['snapshot']['paid']) + sumd(p['amount'] for p in x['events'] if p['id'] not in x['snapshot']['included_events'] and p['date'] > x['snapshot']['as_of'])
            check(key, paid, D(e['paid']))
        elif i == 9:
            check(key + ':total', D(x['trade_ap']) + D(x['loan']), D(e['ledger_total']))
            check(key + ':comparable', D(x['trade_ap']) - D(x['invoice_ap']), D(e['difference']))
        elif i == 10:
            check(key + ':distinct', len({r['id'] for r in x['payments']}), e['real_payment_count'])
            check(key + ':sum', sumd(r['amount'] for r in x['payments']), D(e['paid']))
            check(key + ':duplicate', x['files'][0]['hash'], x['files'][1]['hash'])
            check(key + ':overlap', e['overlap_state'], 'excluded_overlap')
        elif i == 11:
            check(key + ':capacity_bound', D(e['max_active_amount']), min(D(x['bank_capacity']), D(x['invoice_capacity'])))
            check(key + ':concurrent execution pending', s['application_execution'], 'not_run')
        elif i == 12:
            check(key + ':today', D(x['gross']) - D(x['paid_total_today']), D(e['current_remaining']))
            check(key + ':missing', [e['overdue'], e['historical_remaining']], [None, None])
        elif i == 13:
            check(key, D(x['plan_total']) - D(x['paid']), D(e['future_payment']))
        elif i == 14:
            check(key + ':missing_tax', x['rows'][0]['tax'], e['tax'])
            check(key + ':unresolved', e['financial_effect'], None)
        elif i == 15:
            check(key + ':undefined_pct', [v['percent'] for v in e['variants']], [None, None])
            check(key + ':zero_delta', D(x['current']) - D(x['previous_variants'][1]), D(e['variants'][1]['delta']))


def edges_capacity_and_coverage():
    baseline = ROOT / 'inputs/baseline'
    check('baseline file count', len(list(baseline.iterdir())), 8)
    check('exact file duplicate', hashlib.sha256((baseline / 'bank.xlsx').read_bytes()).hexdigest(), hashlib.sha256((ROOT / 'inputs/edges/exact-duplicate-bank.xlsx').read_bytes()).hexdigest())
    weird = load_workbook(ROOT / 'inputs/edges/weird-headers.xlsx', read_only=True)
    check('leading zero account', weird.active['B2'].value, '0000123')
    weird.close()
    count = 0
    amounts = []
    identifiers = []
    for f in sorted((ROOT / 'inputs/capacity').glob('*.csv')):
        rows = crows(str(f.relative_to(ROOT)))
        check('capacity file size:' + f.name, len(rows), 2000)
        count += len(rows)
        amounts.extend(D(r['amount']) for r in rows)
        identifiers.extend(r['reference'] for r in rows)
    check('capacity row total', count, 20000)
    check('capacity unique refs', len(set(identifiers)), 20000)
    check('capacity net', sum(amounts), D(0))
    check('capacity inflow', sum(a for a in amounts if a > 0), D('100000'))
    check('capacity outflow', -sum(a for a in amounts if a < 0), D('100000'))
    for i in range(1, 15):
        s = read_json(f'scenarios/T{i:02}.json')
        check('technical scenario has assertions:' + s['id'], bool(s['assertions']), True)
        check('technical not run:' + s['id'], s['execution_status'], 'specified_not_executed')
    errors = xrows('inputs/edges/invalid-rows.xlsx', 'خطاها')
    invalid = {}
    for row in errors:
        try:
            date.fromisoformat(row['date'])
        except ValueError:
            invalid[row['row_id']] = 'invalid_date'
            continue
        try:
            value = D(str(row['amount']))
        except Exception:
            invalid[row['row_id']] = 'invalid_amount'
            continue
        if -value.as_tuple().exponent > 6:
            invalid[row['row_id']] = 'excess_precision'
    check('invalid row classification', invalid, read_json('scenarios/edge-expectations.json')['invalid-rows.xlsx']['invalid_rows'])
    try:
        with (ROOT / 'inputs/edges/malformed.csv').open() as f:
            list(csv.reader(f, strict=True))
    except csv.Error:
        check('malformed CSV', True, True)
    else:
        raise AssertionError('Malformed CSV unexpectedly valid')
    for f in ROOT.rglob('*.xlsx'):
        workbook = load_workbook(f, read_only=True, data_only=False)
        for sheet in workbook:
            check('no executable formula:' + f.name + ':' + sheet.title, any(c.data_type == 'f' for row in sheet for c in row), False)
        workbook.close()


def reference_contracts_and_hashes():
    manifest = read_json('expected/source-manifest.json')
    for item in manifest['files']:
        p = ROOT / item['path']
        check('source hash:' + p.name, hashlib.sha256(p.read_bytes()).hexdigest(), item['sha256'])
        check('source bytes:' + p.name, p.stat().st_size, item['bytes'])
    original_manifest = {k: manifest[k] for k in ('schema_version', 'company_id', 'files')}
    fingerprint = hashlib.sha256(json.dumps(original_manifest, sort_keys=True, separators=(',', ':')).encode()).hexdigest()
    check('input fingerprint', fingerprint, manifest['fixture_input_fingerprint'])
    api = read_json('expected/metric-observations.json')
    golden = read_json('expected/baseline.json')
    aliases = read_json('scenarios/identity-aliases.json')['aliases']
    check('API input fingerprint', api['fixture_input_fingerprint'], fingerprint)
    for metric in api['metrics']:
        required = {'metric_key', 'basis', 'scope', 'value', 'unit', 'status', 'as_of_date', 'period_id', 'comparison', 'reason_codes', 'coverage', 'formula_version', 'manifest_id', 'evidence_refs'}
        check('API metric required fields:' + metric['metric_key'], required <= metric.keys(), True)
        check('API amount string:' + metric['metric_key'], isinstance(metric['value'], str), True)
        check('API company isolation:' + metric['metric_key'], metric['scope']['company_id'], aliases['company-a'])
        UUID(metric['manifest_id'])
        UUID(metric['period_id'])
        if metric['metric_key'] in golden['metrics']:
            check('API golden amount:' + metric['metric_key'], metric['value'], golden['metrics'][metric['metric_key']]['current'])
    null_metric = read_json('expected/nullable-metric-example.json')
    check('API null not zero', null_metric['value'], None)
    check('API null reason present', bool(null_metric['reason_codes']), True)
    source = crows('inputs/edges/formula-text.csv')
    safe = crows('expected/formula-text-safe-export.csv')
    for raw, exported in zip(source, safe, strict=True):
        expected = "'" + raw['description'] if raw['description'].startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else raw['description']
        check('safe text export:' + raw['row_id'], exported['description'], expected)
    check('error export row IDs', [r['row_id'] for r in crows('expected/invalid-rows.csv')], ['00002', '00003', '00004'])


if __name__ == '__main__':
    # Sidecars are diagnostics emitted by the authoring tool, not fixture inputs.
    for sidecar in ROOT.rglob('*.inspect.ndjson'):
        sidecar.unlink()
    baseline()
    domain_cases()
    edges_capacity_and_coverage()
    reference_contracts_and_hashes()
    write_json('validation-report.json', {
        'schema_version': 'fixture-1', 'scope': 'saved_fixture_consistency_only',
        'status': 'passed', 'checks_passed': len(CHECKS), 'checks': CHECKS,
        'financial_cases': {'specified': 15, 'reference_checks': 'passed', 'application_engine': 'not_run'},
        'technical_cases': {'specified': 14, 'application_integration': 'not_run'},
        'accountant_approval': 'pending', 'provider_calls': 0,
        'not_executed': ['application engine', 'database concurrency/RLS', 'model interpretation quality', 'browser and reduced motion', 'five-user load', 'backup/restore', 'XLS legacy parser', 'macro/malware and decompression limits'],
    })
    print(json.dumps({'status': 'passed', 'fixture_checks': len(CHECKS), 'application_tests': 'not_run'}))
